#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_comparacion.py — Entregable 3: Comparación del indicador de Participación
Ciudadana entre (a) baseline 2025 con la FÓRMULA ANTIGUA (legacy 2025, 4 variables),
(b) baseline 2025 con la FÓRMULA NUEVA (CENIA v2, 5 variables, sobre la
recodificación 2026), y (c) el ciclo 2026 (valores del entregable
BBDD_casos_incluidos_ILIA2026.xlsx).

Salida: data/final/Comparacion_2025_2026_ILIA.xlsx — 5 pestañas:
  1 · BBDD 2025 (insumo)            28 casos + recodificación + flags POR FÓRMULA
  2 · Cálculo 2025 · fórmula antigua legacy 4 var, con control vs oficial publicado
  3 · Cálculo 2025 · fórmula nueva   CENIA v2 sobre los mismos 28 casos
  4 · Comparación                    indicador/sub/variables + deltas
  5 · Gráficos                       barras agrupadas (paleta validada CVD-safe)

Todo cálculo es FÓRMULA Excel auditable; los únicos valores pegados como datos son
(i) el oficial 2025 publicado (hoja «Gráfico - Índice» de la BBDD) y (ii) los
resultados 2026 (con fórmulas auditables en su propio archivo).
"""
import csv
import os
import unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIS = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(THIS, ".."))
BBDD = os.path.join(ROOT, "data", "baseline", "BBDD_IA_Participacion_2025.xlsx")
RECOD = os.path.join(ROOT, "data", "baseline", "recodificacion_2025.csv")
D1 = os.path.join(ROOT, "data", "final", "BBDD_casos_incluidos_ILIA2026.xlsx")
OUT = os.path.join(ROOT, "data", "final", "Comparacion_2025_2026_ILIA.xlsx")

PAISES = ["AR","BO","BR","CL","CO","CR","CU","DO","EC","GT","HN","JM","MX","PA","PE","PY","SV","TT","UY","VE"]

AZUL_DATO = Font(name="Arial", size=9, color="0000FF")          # dato de entrada
NEGRO = Font(name="Arial", size=9)                               # fórmula
HDR_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FILL2 = PatternFill("solid", fgColor="4472C4")
PARAM_FILL = PatternFill("solid", fgColor="FFFF00")
TITLE_FONT = Font(name="Arial", size=11, bold=True)
NOTA_FONT = Font(name="Arial", size=8, italic=True)
BOLD9 = Font(name="Arial", size=9, bold=True)
THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Paleta categórica validada (dataviz · validate_palette.js: ALL CHECKS PASS)
COLOR_2025_ANTIGUA = "2A78D6"   # azul   — serie 1
COLOR_2025_NUEVA = "008300"     # verde  — serie 2
COLOR_2026 = "E87BA4"           # magenta — serie 3

# --------------------------------------------------------------------------
# 1. Leer insumos
# --------------------------------------------------------------------------
wb_b = load_workbook(BBDD, data_only=True)
ws_b = wb_b["BBDD Casos"]
casos = []
for row in ws_b.iter_rows(min_row=2, values_only=True):
    if not row[0] or not row[1] or not str(row[0]).strip() or not str(row[1]).strip():
        continue
    casos.append({
        "pais": str(row[0]).strip(), "caso": str(row[1]).strip(), "ano": row[2],
        "desc": str(row[4] or "").strip(), "org": str(row[5] or "").strip(),
        "tipo_org": str(row[6] or "").strip(), "tipos_ia": str(row[9] or "").strip(),
        "proceso": str(row[11] or "").strip(), "etapa": str(row[12] or "").strip(),
        "dev": str(row[14] or "").strip(), "origen": str(row[15] or "").strip(),
        "cont": str(row[17] or "").strip(),
    })
assert len(casos) == 28, f"BBDD 2025: se esperaban 28 casos, hay {len(casos)}"

recod = {}
with open(RECOD, newline="", encoding="utf-8") as fh:
    for r in csv.DictReader(fh):
        ov = (r.get("nivel_consolidacion_override") or "").strip()
        niv = ov if ov else (r.get("nivel_consolidacion") or "").strip()
        conv = (r.get("tipo_convocante") or "").replace("|", "; ")
        recod[(r["pais"].strip(), r["caso"].strip().lower())] = {
            "conv": conv, "nivel": int(niv) if niv else None,
            "nota": (r.get("nota") or "").strip(),
        }
for c in casos:
    r = recod.get((c["pais"], c["caso"].lower()))
    assert r is not None, f"sin recodificación: {c['pais']} {c['caso']}"
    c["conv"], c["nivel"], c["nota_recod"] = r["conv"], r["nivel"], r["nota"]

# Oficial 2025 publicado (hoja «Gráfico - Índice»; TT no estaba en el universo 2025)
oficial = {}
for row in wb_b["Gráfico - Índice"].iter_rows(min_row=2, values_only=True):
    if row[0]:
        oficial[str(row[0]).strip()] = (row[1], row[2], row[3])

# Resultados 2026 (valores cacheados del entregable 1, ya recalculado y verificado)
wb1 = load_workbook(D1, data_only=True)
w26 = wb1["2 · Índice por país"]
h26 = {w26.cell(2, c).value: c for c in range(1, w26.max_column + 1)}
res26 = {}
for r in range(3, 23):
    p = w26.cell(r, 1).value
    if p in PAISES:
        res26[p] = {
            "v1": w26.cell(r, h26["V1 Tipos de proceso"]).value,
            "v2": w26.cell(r, h26["V2 Etapas de uso"]).value,
            "v3": w26.cell(r, h26["V3 Continuidad"]).value,
            "v4": w26.cell(r, h26["V4 Convocante"]).value,
            "v5": w26.cell(r, h26["V5 Cantidad"]).value,
            "s1r": w26.cell(r, h26["Sub1 red."]).value,
            "s2r": w26.cell(r, h26["Sub2 red."]).value,
            "ind": w26.cell(r, h26["Indicador final"]).value,
        }
assert len(res26) == 20 and all(v["ind"] is not None for v in res26.values()), \
    "El entregable 1 no tiene valores cacheados: correr recalc.py primero"

# Tokens de sistemas de IA 2025 (distintos, en minúscula, orden de aparición)
def tokens(s):
    return [t.strip().lower() for t in str(s or "").replace(";", ",").split(",") if t.strip()]

sis_tokens = []
for c in casos:
    for t in tokens(c["tipos_ia"]):
        if t not in sis_tokens:
            sis_tokens.append(t)

wb = Workbook()

# ==========================================================================
# PESTAÑA 1 · BBDD 2025 (insumo)
# ==========================================================================
w1 = wb.active
w1.title = "1 · BBDD 2025 (insumo)"

DATA_COLS = [
    ("País", 6), ("Caso", 30), ("Año", 6), ("Descripción breve", 45),
    ("Organización a cargo", 30), ("Tipo de Organización (2025)", 12),
    ("Tipo de Proceso", 22), ("Etapa de Participación", 18), ("Continuidad (2025)", 12),
    ("Tipo de Desarrollador", 16), ("Origen Desarrollador", 14), ("Tipos de IA (2025)", 22),
    ("Convocante 7-cat (recodificación 2026, BORRADOR)", 26),
    ("Nivel 0-3 (recodificación 2026, override aplicado)", 12),
    ("Nota recodificación", 28),
]
# columnas derivadas (fórmulas)
NORM_DEFS = [  # (título, col fuente)
    ("norm proceso", 7), ("norm etapa", 8), ("norm desarrollador", 10),
    ("norm origen", 11), ("norm sistemas IA", 12), ("norm convocante", 13),
]
TP_FLAGS = [("TP: Participación Digital (0/1)", ["participación digital"]),
            ("TP: Gobernanza Colaborativa (0/1)", ["governanza colaborativa", "gobernanza colaborativa"]),
            ("TP: Mini públicos (0/1)", ["mini públicos"]),
            ("TP: Referendos e Inic. Populares (0/1)", ["referendos e iniciativas populares"]),
            ("TP: Presupuestos Participativos (0/1)", ["presupuestos participativos"])]
ET_FLAGS = [("ET: Planificación (0/1)", "planif"), ("ET: Implementación (0/1)", "implement"),
            ("ET: Análisis (0/1)", "análisis"), ("ET: Traducción Política (0/1)", "traducción")]
ORG_FLAGS = [("ORG: Público (0/1)", "público"), ("ORG: Privado (0/1)", "privado"),
             ("ORG: Mixto (0/1)", "mixto")]
CONT_FLAGS = [("CONT: Uso único (0/1)", "único"), ("CONT: Plataforma (0/1)", "plataforma")]
DEV_FLAGS = [("DEV: Industria (0/1)", "industria"), ("DEV: Academia (0/1)", "academia"),
             ("DEV: Gobierno (0/1)", "gobierno"), ("DEV: Sociedad Civil (0/1)", "sociedad civil")]
ORIG_FLAGS = [("ORIGEN: Nacional (0/1)", "nacional"), ("ORIGEN: Internacional (0/1)", "internacional")]
CONV_FLAGS = [("CONV: gobierno local (0/1)", "gobierno local"),
              ("CONV: gobierno nacional (0/1)", "gobierno nacional"),
              ("CONV: otra institución pública (0/1)", "otra institución pública"),
              ("CONV: empresa (0/1)", "empresa"), ("CONV: universidad (0/1)", "universidad"),
              ("CONV: sociedad civil (0/1)", "sociedad civil"),
              ("CONV: organización internacional (0/1)", "organización internacional")]
CONV_CODES = ["GL", "GN", "OIP", "EMP", "UNI", "SC", "OI"]

col = {}
ci = 1
for name, width in DATA_COLS:
    col[name.split(" (")[0]] = ci
    w1.column_dimensions[get_column_letter(ci)].width = width
    ci += 1
c_norm0 = ci
for name, _src in NORM_DEFS:
    col[name] = ci
    w1.column_dimensions[get_column_letter(ci)].width = 4
    ci += 1
for group in (TP_FLAGS, ET_FLAGS, ORG_FLAGS, CONT_FLAGS, DEV_FLAGS, ORIG_FLAGS, CONV_FLAGS):
    for item in group:
        col[item[0]] = ci
        w1.column_dimensions[get_column_letter(ci)].width = 5
        ci += 1
col["Nivel efectivo"] = ci; w1.column_dimensions[get_column_letter(ci)].width = 9; ci += 1
col["Clave combinación convocantes"] = ci; w1.column_dimensions[get_column_letter(ci)].width = 16; ci += 1
col["Combinación 1ª aparición (país)"] = ci; w1.column_dimensions[get_column_letter(ci)].width = 9; ci += 1
NCOL1 = ci - 1

hdrs = ([n for n, _ in DATA_COLS] + [n for n, _ in NORM_DEFS]
        + [n for n, _ in TP_FLAGS] + [n for n, _ in ET_FLAGS] + [n for n, _ in ORG_FLAGS]
        + [n for n, _ in CONT_FLAGS] + [n for n, _ in DEV_FLAGS] + [n for n, _ in ORIG_FLAGS]
        + [n for n, _ in CONV_FLAGS]
        + ["Nivel efectivo", "Clave combinación convocantes", "Combinación 1ª aparición (país)"])
for c, h in enumerate(hdrs, start=1):
    cell = w1.cell(1, c, h)
    cell.font = HDR_FONT
    cell.fill = HDR_FILL if c <= len(DATA_COLS) else HDR_FILL2
    cell.alignment = CENTER
    cell.border = THIN
w1.row_dimensions[1].height = 52

def L(name):
    return get_column_letter(col[name])

for i, c in enumerate(casos):
    r = i + 2
    vals = [c["pais"], c["caso"], c["ano"], c["desc"], c["org"], c["tipo_org"],
            c["proceso"], c["etapa"], c["cont"], c["dev"], c["origen"], c["tipos_ia"],
            c["conv"], c["nivel"], c["nota_recod"]]
    for j, v in enumerate(vals, start=1):
        cell = w1.cell(r, j, v)
        cell.font = AZUL_DATO
        cell.border = THIN
        if j in (4, 5, 7, 8, 12, 13, 15):
            cell.alignment = WRAP
    # normalizadas
    for (name, src), _ in zip(NORM_DEFS, range(6)):
        f = (f'=";"&LOWER(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('
             f'{get_column_letter(src)}{r},"; ",";"),", ",";"),",",";"))&";"')
        w1.cell(r, col[name], f).font = NEGRO
    # flags TP (tokens exactos, con las dos grafías de Governanza/Gobernanza)
    for name, toks in TP_FLAGS:
        n = L("norm proceso")
        conds = ",".join(f'ISNUMBER(SEARCH(";{t};",{n}{r}))' for t in toks)
        f = f"=IF(OR({conds}),1,0)" if len(toks) > 1 else f"=IF({conds},1,0)"
        w1.cell(r, col[name], f).font = NEGRO
    # flags etapa (subcadena: cubre «Traducción» y «Traducción Política»)
    for name, sub in ET_FLAGS:
        w1.cell(r, col[name], f'=IF(ISNUMBER(SEARCH("{sub}",{L("norm etapa")}{r})),1,0)').font = NEGRO
    # flags org 2025 (celda univaluada)
    for name, sub in ORG_FLAGS:
        w1.cell(r, col[name], f'=IF(ISNUMBER(SEARCH("{sub}",{get_column_letter(6)}{r})),1,0)').font = NEGRO
    # flags continuidad 2025
    for name, sub in CONT_FLAGS:
        w1.cell(r, col[name], f'=IF(ISNUMBER(SEARCH("{sub}",{get_column_letter(9)}{r})),1,0)').font = NEGRO
    # flags desarrollador / origen / convocante (token exacto)
    for name, tok in DEV_FLAGS:
        w1.cell(r, col[name], f'=IF(ISNUMBER(SEARCH(";{tok};",{L("norm desarrollador")}{r})),1,0)').font = NEGRO
    for name, tok in ORIG_FLAGS:
        w1.cell(r, col[name], f'=IF(ISNUMBER(SEARCH(";{tok};",{L("norm origen")}{r})),1,0)').font = NEGRO
    for name, tok in CONV_FLAGS:
        w1.cell(r, col[name], f'=IF(ISNUMBER(SEARCH(";{tok};",{L("norm convocante")}{r})),1,0)').font = NEGRO
    # nivel efectivo (caducidad de anuncios; parámetros en pestaña 3)
    w1.cell(r, col["Nivel efectivo"],
            f"=IF(N{r}=\"\",\"\",IF(AND(N{r}=1,('3 · Cálculo 2025 · f. nueva'!$C$25-C{r})>"
            f"'3 · Cálculo 2025 · f. nueva'!$C$26),0,N{r}))").font = NEGRO
    # clave de combinación (orden canónico GL GN OIP EMP UNI SC OI)
    fcols = [L(n) for n, _ in CONV_FLAGS]
    rng = f"{fcols[0]}{r}:{fcols[-1]}{r}"
    parts = "&".join(f'IF({cc}{r}=1,"{code}+","")' for cc, code in zip(fcols, CONV_CODES))
    w1.cell(r, col["Clave combinación convocantes"], f'=IF(SUM({rng})<2,"",{parts})').font = NEGRO
    # 1ª aparición de la combinación en el país (rangos expansivos desde fila 1)
    ck = L("Clave combinación convocantes")
    w1.cell(r, col["Combinación 1ª aparición (país)"],
            f'=IF({ck}{r}="",0,IF(COUNTIFS($A$1:A{r-1},A{r},${ck}$1:{ck}{r-1},{ck}{r})=0,1,0))').font = NEGRO

w1.freeze_panes = "C2"
w1.auto_filter.ref = f"A1:{get_column_letter(NCOL1)}29"

NOTAS1 = [
    "NOTAS Y SUPUESTOS — pestaña 1",
    "· AZUL = dato (28 casos de la BBDD 2025, solo lectura, + recodificación 2026 BORRADOR de data/baseline/recodificacion_2025.csv: "
    "convocante 7-cat y nivel 0-3 con override aplicado). NEGRO = fórmula.",
    "· Las columnas «norm …» normalizan separadores («, » y «;» → «;») y minúsculas; los flags 0/1 buscan el token exacto entre «;» "
    "(imprescindible en Origen: «nacional» es subcadena de «internacional»).",
    "· «Governanza Colaborativa» (grafía 2025) y «Gobernanza Colaborativa» cuentan como el mismo tipo de proceso.",
    "· Nivel efectivo: un anuncio (Nivel=1) caduca a los 2 años sin implementación (parámetros en pestaña 3). "
    "Solo aplica a la fórmula NUEVA; la fórmula antigua usa Continuidad 2025 (Uso único / Plataforma).",
    "· Clave de combinación: conjunto de ≥2 tipos de convocante que co-convocan la misma iniciativa "
    "(GL=gob. local, GN=gob. nacional, OIP=otra inst. pública, EMP=empresa, UNI=universidad, SC=sociedad civil, OI=org. internacional).",
    "· Los flags ORG/CONT alimentan SOLO la fórmula antigua; CONV/Nivel efectivo alimentan SOLO la nueva; "
    "TP/ET/DEV/ORIGEN/sistemas alimentan ambas.",
]
for k, t in enumerate(NOTAS1):
    cell = w1.cell(32 + k, 1, t)
    cell.font = TITLE_FONT if k == 0 else NOTA_FONT
    w1.merge_cells(start_row=32 + k, start_column=1, end_row=32 + k, end_column=14)

# ==========================================================================
# helpers para pestañas de cálculo
# ==========================================================================
P1 = f"'{w1.title}'"

def p1(name, absolute=True):
    letter = L(name)
    return f"{P1}!${letter}$2:${letter}$29" if absolute else f"{P1}!{letter}2:{letter}29"

def style_hdr(ws, r, c, text, fill=HDR_FILL):
    cell = ws.cell(r, c, text)
    cell.font = HDR_FONT
    cell.fill = fill
    cell.alignment = CENTER
    cell.border = THIN
    return cell

def fnum(ws, r, c, formula, fmt="0.0", font=NEGRO, bold=False):
    cell = ws.cell(r, c, formula)
    cell.font = BOLD9 if bold else font
    cell.number_format = fmt
    cell.border = THIN
    return cell

def grilla(ws, titulo, row0, flag_names, total_title="SUMA"):
    """Grilla 20 países × columnas de presencia (0/1) desde los flags de pestaña 1."""
    ws.cell(row0, 1, titulo).font = TITLE_FONT
    style_hdr(ws, row0 + 1, 1, "País", HDR_FILL2)
    for j, fn in enumerate(flag_names, start=2):
        style_hdr(ws, row0 + 1, j, fn.replace(" (0/1)", ""), HDR_FILL2)
    style_hdr(ws, row0 + 1, len(flag_names) + 2, total_title, HDR_FILL2)
    for i, p in enumerate(PAISES):
        r = row0 + 2 + i
        ws.cell(r, 1, p).font = BOLD9
        for j, fn in enumerate(flag_names, start=2):
            fnum(ws, r, j, f'=IF(COUNTIFS({p1("País")},$A{r},{p1(fn)},1)>0,1,0)', "0")
        cletters = f"B{r}:{get_column_letter(len(flag_names)+1)}{r}"
        fnum(ws, r, len(flag_names) + 2, f"=SUM({cletters})", "0", bold=True)
    return row0 + 2, len(flag_names) + 2  # primera fila de datos, columna SUMA

# ==========================================================================
# PESTAÑA 2 · Cálculo 2025 · fórmula antigua
# ==========================================================================
w2 = wb.create_sheet("2 · Cálculo 2025 · f. antigua")
w2.cell(1, 1, "CÁLCULO 2025 CON LA FÓRMULA ANTIGUA (legacy 2025 · 4 variables · como se publicó el ILIA 2025)").font = TITLE_FONT
H2 = ["País", "N casos", "nº tipos de proceso", "V Tipos (nº÷5×100)", "nº etapas",
      "V Etapas (nº÷4×100)", "nº modalidades continuidad (UU/PL)", "V Continuidad (nº÷2×100)",
      "nº tipos org (Púb/Priv/Mixto)", "V Público-Privado (nº÷3×100)", "Sub1 Uso (promedio 4 V)",
      "Sub1 red.", "dev. nacionales (tipos)", "dev. internacionales (tipos)", "V Desarrollador",
      "nº sistemas IA", "V Tipos de IA", "Sub2 Desarrollo", "Sub2 red.",
      "INDICADOR 2025 · fórmula antigua", "Sub1 oficial 2025", "Sub2 oficial 2025",
      "Indicador oficial 2025", "Control |Δ|≤1"]
for c, h in enumerate(H2, start=1):
    style_hdr(w2, 2, c, h)
    w2.column_dimensions[get_column_letter(c)].width = 11
w2.column_dimensions["A"].width = 6
w2.row_dimensions[2].height = 58

R0 = 3   # primera fila de países en la tabla principal
# posiciones de grillas (abajo)
G_TIPOS, G_ET, G_CONT, G_ORG, G_DNAC, G_DINT, G_SIS = 27, 51, 75, 99, 123, 147, 171

tp_names = [n for n, _ in TP_FLAGS]
et_names = [n for n, _ in ET_FLAGS]
org_names = [n for n, _ in ORG_FLAGS]
cont_names = [n for n, _ in CONT_FLAGS]
dev_names = [n for n, _ in DEV_FLAGS]

for i, p in enumerate(PAISES):
    r = R0 + i
    w2.cell(r, 1, p).font = BOLD9
    w2.cell(r, 1).border = THIN
    fnum(w2, r, 2, f'=COUNTIF({p1("País")},$A{r})', "0")
    fnum(w2, r, 3, f"=${get_column_letter(len(tp_names)+2)}${G_TIPOS+2+i}", "0")
    fnum(w2, r, 4, f"=C{r}/5*100")
    fnum(w2, r, 5, f"=${get_column_letter(len(et_names)+2)}${G_ET+2+i}", "0")
    fnum(w2, r, 6, f"=E{r}/4*100")
    fnum(w2, r, 7, f"=${get_column_letter(len(cont_names)+2)}${G_CONT+2+i}", "0")
    fnum(w2, r, 8, f"=G{r}/2*100")
    fnum(w2, r, 9, f"=${get_column_letter(len(org_names)+2)}${G_ORG+2+i}", "0")
    fnum(w2, r, 10, f"=I{r}/3*100")
    fnum(w2, r, 11, f"=(D{r}+F{r}+H{r}+J{r})/4")
    fnum(w2, r, 12, f"=ROUND(K{r},0)", "0", bold=True)
    fnum(w2, r, 13, f"=${get_column_letter(len(dev_names)+2)}${G_DNAC+2+i}", "0")
    fnum(w2, r, 14, f"=${get_column_letter(len(dev_names)+2)}${G_DINT+2+i}", "0")
    fnum(w2, r, 15, f"=(IF($C$24=0,0,M{r}/$C$24)+IF($C$25=0,0,N{r}/$C$25))/2*100")
    fnum(w2, r, 16, f"=${get_column_letter(len(sis_tokens)+2)}${G_SIS+2+i}", "0")
    fnum(w2, r, 17, f"=IF($C$26=0,0,P{r}/$C$26*100)")
    fnum(w2, r, 18, f"=(O{r}+Q{r})/2")
    fnum(w2, r, 19, f"=ROUND(R{r},0)", "0", bold=True)
    fnum(w2, r, 20, f"=ROUND((L{r}+S{r})/2,0)", "0", bold=True)
    w2.cell(r, 20).fill = PatternFill("solid", fgColor="DCE6F1")
    of = oficial.get(p)
    for j, v in zip((21, 22, 23), of if of else (None, None, None)):
        cell = w2.cell(r, j, v if v is not None else "s/d")
        cell.font = AZUL_DATO
        cell.number_format = "0"
        cell.border = THIN
    fnum(w2, r, 24, f'=IF(W{r}="s/d","s/d",IF(ABS(T{r}-W{r})<=1,"✓","✗"))', "General", bold=True)

# parámetros (máximos relativos del ciclo 2025, por fórmula)
w2.cell(23, 1, "MÁXIMOS RELATIVOS DEL CICLO 2025 (fórmulas MAX sobre la tabla)").font = TITLE_FONT
for rr, (lbl, f) in enumerate([
        ("Máx. desarrolladores nacionales", f"=MAX(M{R0}:M{R0+19})"),
        ("Máx. desarrolladores internacionales", f"=MAX(N{R0}:N{R0+19})"),
        ("Máx. sistemas de IA", f"=MAX(P{R0}:P{R0+19})")], start=24):
    w2.cell(rr, 1, lbl).font = NOTA_FONT
    fnum(w2, rr, 3, f, "0", bold=True).fill = PARAM_FILL

w2.cell(26, 5, "Fuente del oficial 2025: hoja «Gráfico - Índice» de data/baseline/BBDD_IA_Participacion_2025.xlsx (dato pegado). "
               "TT = s/d (no estaba en el universo 2025). Diferencias de ±1 provienen del redondeo del motor original (banker's) "
               "vs ROUND de Excel (half-up).").font = NOTA_FONT
w2.merge_cells(start_row=26, start_column=5, end_row=26, end_column=20)

grilla(w2, "GRILLA — Tipos de proceso presentes por país (20×5)", G_TIPOS, tp_names)
grilla(w2, "GRILLA — Etapas presentes por país (20×4)", G_ET, et_names)
grilla(w2, "GRILLA — Modalidades de continuidad 2025 (Uso único / Plataforma) (20×2)", G_CONT, cont_names)
grilla(w2, "GRILLA — Tipos de organización 2025 (Público/Privado/Mixto) (20×3)", G_ORG, org_names)

# grillas de desarrollador × origen (nacional / internacional)
for row0, titulo, orig_flag in [
        (G_DNAC, "GRILLA — Tipos de desarrollador con origen NACIONAL (20×4)", "ORIGEN: Nacional (0/1)"),
        (G_DINT, "GRILLA — Tipos de desarrollador con origen INTERNACIONAL (20×4)", "ORIGEN: Internacional (0/1)")]:
    w2.cell(row0, 1, titulo).font = TITLE_FONT
    style_hdr(w2, row0 + 1, 1, "País", HDR_FILL2)
    for j, fn in enumerate(dev_names, start=2):
        style_hdr(w2, row0 + 1, j, fn.replace(" (0/1)", ""), HDR_FILL2)
    style_hdr(w2, row0 + 1, len(dev_names) + 2, "SUMA", HDR_FILL2)
    for i, p in enumerate(PAISES):
        r = row0 + 2 + i
        w2.cell(r, 1, p).font = BOLD9
        for j, fn in enumerate(dev_names, start=2):
            fnum(w2, r, j,
                 f'=IF(COUNTIFS({p1("País")},$A{r},{p1(fn)},1,{p1(orig_flag)},1)>0,1,0)', "0")
        fnum(w2, r, len(dev_names) + 2, f"=SUM(B{r}:{get_column_letter(len(dev_names)+1)}{r})", "0", bold=True)

# grilla de sistemas de IA 2025 (tokens exactos)
w2.cell(G_SIS, 1, f"GRILLA — Sistemas de IA distintos por país ({len(sis_tokens)} tokens 2025; comparación textual exacta)").font = TITLE_FONT
style_hdr(w2, G_SIS + 1, 1, "País", HDR_FILL2)
for j, t in enumerate(sis_tokens, start=2):
    style_hdr(w2, G_SIS + 1, j, t, HDR_FILL2)
    w2.column_dimensions[get_column_letter(j)].width = 11
style_hdr(w2, G_SIS + 1, len(sis_tokens) + 2, "SUMA", HDR_FILL2)
for i, p in enumerate(PAISES):
    r = G_SIS + 2 + i
    w2.cell(r, 1, p).font = BOLD9
    for j, t in enumerate(sis_tokens, start=2):
        fnum(w2, r, j, f'=IF(COUNTIFS({p1("País")},$A{r},{p1("norm sistemas IA")},"*;{t};*")>0,1,0)', "0")
    fnum(w2, r, len(sis_tokens) + 2, f"=SUM(B{r}:{get_column_letter(len(sis_tokens)+1)}{r})", "0", bold=True)

w2.freeze_panes = "B3"

# ==========================================================================
# PESTAÑA 3 · Cálculo 2025 · fórmula nueva (CENIA v2)
# ==========================================================================
w3 = wb.create_sheet("3 · Cálculo 2025 · f. nueva")
w3.cell(1, 1, "CÁLCULO 2025 CON LA FÓRMULA NUEVA (CENIA v2 · 5 variables · misma base 2025 recodificada)").font = TITLE_FONT
H3 = ["País", "N iniciativas", "nº tipos", "V1 Tipos de proceso", "nº etapas", "V2 Etapas de uso",
      "nivel máx (efectivo)", "V3 Continuidad", "nº tipos gub", "nº tipos no-gub",
      "nº combinaciones", "V4 Convocante", "V5 Cantidad", "Sub1 Uso", "Sub1 red.",
      "V Desarrollador (=f. antigua)", "V Tipos de IA (=f. antigua)", "Sub2 Desarrollo (=f. antigua)",
      "Sub2 red.", "INDICADOR 2025 · fórmula nueva"]
for c, h in enumerate(H3, start=1):
    style_hdr(w3, 2, c, h)
    w3.column_dimensions[get_column_letter(c)].width = 11
w3.column_dimensions["A"].width = 6
w3.row_dimensions[2].height = 58

G_GUB, G_NOGUB = 45, 69
gub_names = [n for n, _ in CONV_FLAGS[:3]]
nogub_names = [n for n, _ in CONV_FLAGS[3:]]
S2 = f"'{w2.title}'"

for i, p in enumerate(PAISES):
    r = R0 + i
    w3.cell(r, 1, p).font = BOLD9
    w3.cell(r, 1).border = THIN
    fnum(w3, r, 2, f'=COUNTIF({p1("País")},$A{r})', "0")
    fnum(w3, r, 3, f"={S2}!C{r}", "0")           # mismos tipos (5) que la antigua
    fnum(w3, r, 4, f"=C{r}/5*100")
    fnum(w3, r, 5, f"={S2}!E{r}", "0")           # mismas etapas (4)
    fnum(w3, r, 6, f"=E{r}/4*100")
    fnum(w3, r, 7, f'=_xlfn.MAXIFS({p1("Nivel efectivo")},{p1("País")},$A{r})', "0")
    fnum(w3, r, 8, f"=G{r}/3*100")
    fnum(w3, r, 9, f"=${get_column_letter(len(gub_names)+2)}${G_GUB+2+i}", "0")
    fnum(w3, r, 10, f"=${get_column_letter(len(nogub_names)+2)}${G_NOGUB+2+i}", "0")
    fnum(w3, r, 11, f'=SUMIFS({p1("Combinación 1ª aparición (país)")},{p1("País")},$A{r})', "0")
    fnum(w3, r, 12, f"=(I{r}/3+J{r}/4+IF($C$27=0,0,K{r}/$C$27))/3*100")
    fnum(w3, r, 13, f"=IF(B{r}=0,$F$29,IF(B{r}<=3,$F$30,IF(B{r}<=6,$F$31,IF(B{r}<=9,$F$32,$F$33))))", "0")
    fnum(w3, r, 14, f"=(D{r}+F{r}+H{r}+L{r}+M{r})/5")
    fnum(w3, r, 15, f"=ROUND(N{r},0)", "0", bold=True)
    fnum(w3, r, 16, f"={S2}!O{r}")
    fnum(w3, r, 17, f"={S2}!Q{r}")
    fnum(w3, r, 18, f"={S2}!R{r}")
    fnum(w3, r, 19, f"=ROUND(R{r},0)", "0", bold=True)
    fnum(w3, r, 20, f"=ROUND((O{r}+S{r})/2,0)", "0", bold=True)
    w3.cell(r, 20).fill = PatternFill("solid", fgColor="E2EFDA")

w3.cell(24, 1, "PARÁMETROS (azul sobre amarillo = entrada; el máximo es fórmula)").font = TITLE_FONT
w3.cell(25, 1, "Año de referencia del ciclo").font = NOTA_FONT
c = w3.cell(25, 3, 2026); c.font = AZUL_DATO; c.fill = PARAM_FILL; c.border = THIN
w3.cell(26, 1, "Caducidad de anuncio (años)").font = NOTA_FONT
c = w3.cell(26, 3, 2); c.font = AZUL_DATO; c.fill = PARAM_FILL; c.border = THIN
w3.cell(27, 1, "Máx. combinaciones del ciclo 2025").font = NOTA_FONT
fnum(w3, 27, 3, f"=MAX(K{R0}:K{R0+19})", "0", bold=True).fill = PARAM_FILL
w3.cell(28, 1, "Umbrales de Cantidad (nº iniciativas → puntaje)").font = TITLE_FONT
for k, (lbl, v) in enumerate([("0", 0), ("1-3", 25), ("4-6", 50), ("7-9", 75), ("10 o más", 100)]):
    w3.cell(29 + k, 5, lbl).font = NOTA_FONT
    c = w3.cell(29 + k, 6, v); c.font = AZUL_DATO; c.fill = PARAM_FILL; c.border = THIN; c.number_format = "0"
w3.cell(35, 1, "Sub2 es idéntico en ambas fórmulas (la metodología 2026 no lo modificó): las columnas P/Q/R referencian la pestaña 2.").font = NOTA_FONT
w3.merge_cells(start_row=35, start_column=1, end_row=35, end_column=14)
w3.cell(36, 1, "V1/V2 (tipos y etapas) usan las MISMAS grillas de la pestaña 2: la definición de esas dos variables no cambió; "
               "lo nuevo es Continuidad por nivel máximo, Convocante (3 subcomponentes con co-convocatoria) y Cantidad.").font = NOTA_FONT
w3.merge_cells(start_row=36, start_column=1, end_row=36, end_column=14)

grilla(w3, "GRILLA — Convocante gubernamental (recodificación 2026) (20×3)", G_GUB, gub_names)
grilla(w3, "GRILLA — Convocante no gubernamental (recodificación 2026) (20×4)", G_NOGUB, nogub_names)
w3.freeze_panes = "B3"

# ==========================================================================
# PESTAÑA 4 · Comparación
# ==========================================================================
w4 = wb.create_sheet("4 · Comparación")
w4.cell(1, 1, "COMPARACIÓN — Indicador de IA en Participación Ciudadana").font = TITLE_FONT
w4.cell(2, 1, "2025 con fórmula antigua y con fórmula nueva = fórmulas que referencian las pestañas 2 y 3 (auditables aquí). "
              "2026 = valores del entregable BBDD_casos_incluidos_ILIA2026.xlsx · pestaña «2 · Índice por país» "
              "(pegados como dato el 2026-07-19; las fórmulas auditables están en ese archivo).").font = NOTA_FONT
w4.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
w4.cell(3, 1, "Δ efecto fórmula = misma base 2025, fórmula nueva − antigua · Δ real = 2026 − 2025 (ambos con fórmula nueva) · "
              "Δ total = 2026 (f. nueva) − 2025 (f. antigua publicada)").font = NOTA_FONT
w4.merge_cells(start_row=3, start_column=1, end_row=3, end_column=12)

H4 = ["País", "Indicador 2025 · fórmula antigua", "Indicador 2025 · fórmula nueva", "Indicador 2026",
      "Δ efecto fórmula (2025n−2025a)", "Δ real (2026−2025n)", "Δ total (2026−2025a)",
      "Sub1 2025 · f. antigua", "Sub1 2025 · f. nueva", "Sub1 2026",
      "Sub2 2025 (ambas fórmulas)", "Sub2 2026",
      "V1 Tipos 2025n", "V1 Tipos 2026", "V2 Etapas 2025n", "V2 Etapas 2026",
      "V3 Continuidad 2025n", "V3 Continuidad 2026", "V4 Convocante 2025n", "V4 Convocante 2026",
      "V5 Cantidad 2025n", "V5 Cantidad 2026"]
HR4 = 5
for c, h in enumerate(H4, start=1):
    style_hdr(w4, HR4, c, h)
    w4.column_dimensions[get_column_letter(c)].width = 11
w4.column_dimensions["A"].width = 6
w4.row_dimensions[HR4].height = 58
S3 = f"'{w3.title}'"

for i, p in enumerate(PAISES):
    r = HR4 + 1 + i
    w4.cell(r, 1, p).font = BOLD9
    w4.cell(r, 1).border = THIN
    fnum(w4, r, 2, f"={S2}!T{R0+i}", "0", bold=True)
    fnum(w4, r, 3, f"={S3}!T{R0+i}", "0", bold=True)
    d = res26[p]
    c26 = w4.cell(r, 4, d["ind"]); c26.font = Font(name="Arial", size=9, bold=True, color="0000FF")
    c26.number_format = "0"; c26.border = THIN
    c26.fill = PatternFill("solid", fgColor="FCE4EC")
    fnum(w4, r, 5, f"=C{r}-B{r}", "+0;-0;0")
    fnum(w4, r, 6, f"=D{r}-C{r}", "+0;-0;0")
    fnum(w4, r, 7, f"=D{r}-B{r}", "+0;-0;0")
    fnum(w4, r, 8, f"={S2}!L{R0+i}", "0")
    fnum(w4, r, 9, f"={S3}!O{R0+i}", "0")
    cc = w4.cell(r, 10, d["s1r"]); cc.font = AZUL_DATO; cc.number_format = "0"; cc.border = THIN
    fnum(w4, r, 11, f"={S2}!S{R0+i}", "0")
    cc = w4.cell(r, 12, d["s2r"]); cc.font = AZUL_DATO; cc.number_format = "0"; cc.border = THIN
    for j, (k25, k26) in enumerate([("D", "v1"), ("F", "v2"), ("H", "v3"), ("L", "v4"), ("M", "v5")]):
        fnum(w4, r, 13 + 2 * j, f"={S3}!{k25}{R0+i}")
        cc = w4.cell(r, 14 + 2 * j, round(float(d[k26]), 1))
        cc.font = AZUL_DATO; cc.number_format = "0.0"; cc.border = THIN
w4.freeze_panes = f"B{HR4+1}"
w4.auto_filter.ref = f"A{HR4}:{get_column_letter(len(H4))}{HR4+20}"

# ==========================================================================
# PESTAÑA 5 · Gráficos
# ==========================================================================
w5 = wb.create_sheet("5 · Gráficos")
w5.cell(1, 1, "GRÁFICOS — todos leen los rangos de la pestaña «4 · Comparación» (recalculan al cambiar los datos)").font = TITLE_FONT

def bar_chart(title, cols, colors, anchor, y_title="Puntaje (0-100)"):
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.title = title
    ch.y_axis.title = y_title
    ch.x_axis.title = "País"
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    ch.height, ch.width = 9.2, 21
    ch.gapWidth = 120
    ch.legend.position = "b"
    cats = Reference(w4, min_col=1, min_row=HR4 + 1, max_row=HR4 + 20)
    for cidx, color in zip(cols, colors):
        ser = Series(Reference(w4, min_col=cidx, min_row=HR4, max_row=HR4 + 20), title_from_data=True)
        ser.graphicalProperties = GraphicalProperties(solidFill=color)
        ch.series.append(ser)
    ch.set_categories(cats)
    w5.add_chart(ch, anchor)

bar_chart("Indicador de IA en Participación Ciudadana — 2025 (fórmula antigua) vs 2025 (fórmula nueva) vs 2026",
          [2, 3, 4], [COLOR_2025_ANTIGUA, COLOR_2025_NUEVA, COLOR_2026], "A3")
bar_chart("Sub-Indicador de Uso (Sub1) — efecto de la fórmula y cambio 2026",
          [8, 9, 10], [COLOR_2025_ANTIGUA, COLOR_2025_NUEVA, COLOR_2026], "A22")
bar_chart("Sub-Indicador de Desarrollo (Sub2) — idéntico en ambas fórmulas 2025 vs 2026",
          [11, 12], [COLOR_2025_ANTIGUA, COLOR_2026], "A41")
bar_chart("V1 · Tipos de proceso — 2025 (f. nueva) vs 2026", [13, 14],
          [COLOR_2025_NUEVA, COLOR_2026], "A60")
bar_chart("V2 · Etapas de uso — 2025 (f. nueva) vs 2026", [15, 16],
          [COLOR_2025_NUEVA, COLOR_2026], "A79")
bar_chart("V3 · Continuidad (nivel máximo) — 2025 (f. nueva) vs 2026", [17, 18],
          [COLOR_2025_NUEVA, COLOR_2026], "A98")
bar_chart("V4 · Organización convocante — 2025 (f. nueva) vs 2026", [19, 20],
          [COLOR_2025_NUEVA, COLOR_2026], "A117")
bar_chart("V5 · Cantidad de iniciativas — 2025 (f. nueva) vs 2026", [21, 22],
          [COLOR_2025_NUEVA, COLOR_2026], "A136")

w5.cell(2, 1, "Azul = 2025 fórmula antigua · Verde = 2025 fórmula nueva · Magenta = 2026. "
              "V1-V5 existen como variables solo en la fórmula nueva (V1/V2 comparten definición con la antigua). "
              "Paleta verificada para daltonismo (ΔE CVD ≥ 8).").font = NOTA_FONT

wb.save(OUT)
print(f"OK -> {OUT}")
print(f"   casos 2025: {len(casos)} · tokens sistemas 2025: {len(sis_tokens)} · países: {len(PAISES)}")
