#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
analizar_reingreso.py — Análisis de reingreso de los casos EXCLUIDOS en 2025.

Revisa los 49 casos excluidos en 2025 y, a la luz del criterio 2026 y de los
hallazgos de la búsqueda, clasifica cada uno en:
  · STAYS_OUT     — la exclusión sigue vigente (atención, no-participación, sin IA…).
  · REVISAR       — proceso/uso elegible plausible donde NO se confirmó IA: vale
                    re-evaluar (entra a la planilla del equipo).
  · YA_REINGRESO  — superado por un caso NUEVO ya capturado (no re-votar aquí).

Salidas: data/gates/reingreso_excluidos.md (análisis completo) y
         data/candidatos/excluidos_revisar.csv (los REVISAR, para la planilla).
"""
from __future__ import annotations
import csv
import os
import sys
import unicodedata

import openpyxl

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))


def _norm(s):
    s = "".join(c for c in unicodedata.normalize("NFKD", str(s)) if not unicodedata.combining(c))
    return s.lower().strip()


# Casos que vuelven a REVISIÓN (proceso elegible donde la IA sobre aportes no se
# confirmó en 2025): clave = subcadena del título; valor = por qué revisar.
REVISAR = {
    "ilabs": "Se describe como integración de analítica de datos, IA y deliberación "
             "ciudadana (mini-públicos, jurados, asambleas) — tipo de proceso ELEGIBLE; "
             "confirmar si la IA procesa el CONTENIDO de los aportes.",
    "que veas": "Bots que recopilan quejas ciudadanas y predicen riesgos en obras públicas "
                "(análogo a dIAra, que sí entra); confirmar IA sobre el contenido de las quejas.",
    "ucampus": "Plataforma oficial del proceso constituyente 2023 (proceso participativo "
               "ELEGIBLE); confirmar si hubo análisis NLP de los aportes (como en cabildos 2016).",
    "asamblea ciudadana itinerante": "Mini-público recurrente del Concejo de Bogotá (proceso "
               "ELEGIBLE); confirmar si incorpora IA sobre el contenido de los aportes.",
    "procuraduria": "Mini-público / diálogo social (proceso ELEGIBLE); confirmar uso de IA "
                    "sobre el contenido de los aportes.",
}
# Casos ya superados por un caso NUEVO capturado (no re-votar aquí).
YA_REINGRESO = {
    "participa+ brasil": "Reingresó como caso NUEVO 'Brasil Participativo' (la plataforma "
                         "federal ahora aplica IA/clustering a las propuestas).",
    "ciudadano a senador": "Concepto ya desplegado: ver caso NUEVO 'e-Cidadania — IA matching "
                           "de ideas legislativas' (Senado Federal, 2026).",
}


def _clasificar(titulo):
    t = _norm(titulo)
    for k, v in YA_REINGRESO.items():
        if _norm(k) in t:
            return "YA_REINGRESO", v
    for k, v in REVISAR.items():
        if k in t:
            return "REVISAR", v
    return "STAYS_OUT", ""


def leer_excluidos():
    wb = openpyxl.load_workbook(os.path.join(ROOT, "data/baseline/BBDD_IA_Participacion_2025.xlsx"), data_only=True)
    ws = wb["BBDD Casos Excluidos"]
    out = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[1] and not row[2]:
            continue
        links = [str(x).strip() for x in (row[5], row[6], row[7]) if x and str(x).strip()]
        out.append({
            "motivo": str(row[0]).strip() if row[0] else "",
            "pais": str(row[1]).strip() if row[1] else "",
            "caso": str(row[2]).strip() if row[2] else "",
            "desc": str(row[3]).strip() if row[3] else "",
            "enlace": links[0] if links else "",
        })
    return out


def main(argv=None):
    ex = leer_excluidos()
    for e in ex:
        e["disp"], e["porque"] = _clasificar(e["caso"])

    revisar = [e for e in ex if e["disp"] == "REVISAR"]
    reentro = [e for e in ex if e["disp"] == "YA_REINGRESO"]
    fuera = [e for e in ex if e["disp"] == "STAYS_OUT"]

    # CSV con los REVISAR (para la planilla)
    csv_path = os.path.join(ROOT, "data/candidatos/excluidos_revisar.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["pais", "caso", "motivo_2025", "por_que_revisar", "enlace"])
        for e in revisar:
            w.writerow([e["pais"], e["caso"], e["motivo"], e["porque"], e["enlace"]])

    # Reporte md
    L = ["# Reingreso de excluidos 2025 — análisis", "",
         "**Pipeline ILIA 2026.** Revisión de los 49 casos excluidos en 2025 bajo el "
         "criterio 2026 y los hallazgos de la búsqueda.", "",
         f"- **A REVISAR (reingreso posible): {len(revisar)}** → van a la planilla del equipo.",
         f"- **Ya reingresaron como caso nuevo: {len(reentro)}** → no re-votar (ya están entre los nuevos).",
         f"- **Exclusión vigente (se mantienen fuera): {len(fuera)}**.", "",
         "## A revisar (reingreso posible)", "",
         "| País | Caso | Motivo 2025 | Por qué revisar |", "|---|---|---|---|"]
    for e in revisar:
        L.append(f"| {e['pais']} | {e['caso']} | {e['motivo']} | {e['porque']} |")
    L += ["", "## Ya reingresaron vía un caso nuevo", "",
          "| País | Caso excluido 2025 | Nota |", "|---|---|---|"]
    for e in reentro:
        L.append(f"| {e['pais']} | {e['caso']} | {e['porque']} |")
    L += ["", "## Exclusión vigente (se mantienen fuera)", "",
          "Mayoría por: atención ciudadana / no es participación / sin IA / civic-tech de "
          "monitoreo / escucha en redes sociales / estudio académico.", "",
          "| País | Caso | Motivo 2025 |", "|---|---|---|"]
    for e in fuera:
        L.append(f"| {e['pais']} | {e['caso']} | {e['motivo']} |")
    L.append("")
    with open(os.path.join(ROOT, "data/gates/reingreso_excluidos.md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))

    print(f"[ok] reingreso: REVISAR={len(revisar)} · YA_REINGRESO={len(reentro)} · STAYS_OUT={len(fuera)}")
    print(f"[ok] {csv_path}")
    print(f"[ok] data/gates/reingreso_excluidos.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
