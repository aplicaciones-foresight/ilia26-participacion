#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_validacion_xlsx.py — Genera el cuaderno de validación manual para el equipo.

Construye `data/gates/validacion_candidatos_ILIA2026.xlsx` a partir de los CSV del
pipeline (candidatos, recodificación 2025, directorio de medios). Pensado para
subir a Google Sheets y trabajar el Gate 1/Gate 2 en equipo: trae la evidencia y
columnas de decisión (con menús desplegables) para que los revisores las llenen.

Reproducible: re-ejecutar tras actualizar los CSV regenera el cuaderno.
"""
from __future__ import annotations
import csv
import datetime
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

HDR = PatternFill("solid", fgColor="1F4E78")          # azul oscuro
HDR_FONT = Font(bold=True, color="FFFFFF")
DEC = PatternFill("solid", fgColor="FFE699")          # ámbar (columnas a llenar)
DEC_FONT = Font(bold=True, color="7F6000")
TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
THIN = Border(*(Side(style="thin", color="D9D9D9"),) * 4)


def _read(path):
    p = os.path.join(ROOT, path)
    if not os.path.exists(p):
        return []
    with open(p, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _hint(notas: str) -> str:
    partes = [p.strip() for p in (notas or "").split(";")]
    drop = ("fuente=", "año baseline", "año=")
    partes = [p for p in partes if p and not any(p.lower().startswith(d) for d in drop)]
    return "; ".join(partes)


def _estilo_header(ws, ncols, dec_from=None):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        if dec_from and c >= dec_from:
            cell.fill = DEC; cell.font = DEC_FONT
        else:
            cell.fill = HDR; cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 30


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# --------------------------------------------------------------------------- #
def hoja_instrucciones(wb, fecha):
    ws = wb.active; ws.title = "Instrucciones"
    ws.sheet_properties.tabColor = "1F4E78"
    lines = [
        ("ILIA 2026 — Validación manual de candidatos (Gate 1 / Gate 2)", TITLE_FONT),
        (f"Generado: {fecha}", None),
        ("", None),
        ("Qué es:", Font(bold=True)),
        ("Planilla para revisar en equipo los casos descubiertos de IA aplicada a participación", None),
        ("ciudadana. Las columnas en ámbar son para que el equipo las complete; el resto es la", None),
        ("evidencia que trajo el pipeline (no editar salvo corrección).", None),
        ("", None),
        ("Criterio de elegibilidad (decisivo):", Font(bold=True)),
        ("Entra solo si, DENTRO de un proceso participativo (consulta, plan de desarrollo,", None),
        ("presupuesto participativo, asamblea/mini-público, iniciativa/referendo, diálogo), se usa", None),
        ("IA sobre el CONTENIDO de los aportes ciudadanos (recolectar, clasificar, analizar/agrupar,", None),
        ("sintetizar, mediar, personalizar).", None),
        ("", None),
        ("Excluir:", Font(bold=True)),
        ("chatbots de atención/FAQ/trámites · conteo/integridad electoral · civic-tech SIN IA ·", None),
        ("votación entre opciones cerradas · consultas SOBRE IA que no usan IA en su proceso ·", None),
        ("encuestas de opinión · papers puramente teóricos sin caso desplegado.", None),
        ("", None),
        ("Cómo llenar (hoja 'Candidatos', columnas ámbar):", Font(bold=True)),
        ("• ¿Elegible?  -> SI / NO / DUDA", None),
        ("• Decisión    -> Aprobar Fase 2 / Verificar / Descartar / Mantener", None),
        ("• Tipo de proceso (2026): Participación Digital, Gobernanza Colaborativa, Mini públicos,", None),
        ("  Referendos e Iniciativas Populares, Presupuestos Participativos (los que apliquen).", None),
        ("• Tipo convocante (2026): gobierno local / gobierno nacional / otra institución pública /", None),
        ("  empresa / universidad / sociedad civil / organización internacional (los que apliquen).", None),
        ("• Nivel consolidación: 0 no hay · 1 anuncio · 2 caso puntual · 3 permanente (recurrencia).", None),
        ("", None),
        ("Leyenda de 'Tipo':", Font(bold=True)),
        ("NUEVO = hallazgo web con evidencia sólida · NUEVO (dudoso) = señal parcial a confirmar ·", None),
        ("baseline 2025 (reverificado) = caso 2025 reconfirmado con actividad 2026 ·", None),
        ("baseline 2025 = caso del índice 2025 · excluido 2025 = se excluyó en 2025 (re-verificar).", None),
        ("", None),
        ("Hojas:", Font(bold=True)),
        ("• Candidatos: la validación principal (98 filas, ordenadas por prioridad).", None),
        ("• Recodificación 2025: confirmar convocante y nivel de los 28 casos baseline.", None),
        ("• Medios (A.10): directorio de medios por país (referencia).", None),
        ("• Resumen: conteos por país y tipo.", None),
    ]
    for i, (txt, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=txt)
        if font: c.font = font
    ws.column_dimensions["A"].width = 100


def hoja_candidatos(wb):
    rows = _read("data/candidatos/candidatos.csv")
    PR = {"nuevo": 1, "nuevo_dudoso": 2}

    def tipo_display(r):
        tb = r["tipo_baseline"]
        if tb == "nuevo": return "NUEVO", 1
        if tb == "nuevo_dudoso": return "NUEVO (dudoso)", 2
        if tb == "activo_2025":
            if "reconfirmado web" in (r["notas"] or ""):
                return "baseline 2025 (reverificado)", 3
            return "baseline 2025", 4
        if tb == "excluido_2025": return "excluido 2025", 6
        return "(sin baseline)", 5

    enr = []
    for r in rows:
        disp, pr = tipo_display(r)
        if r["tipo_baseline"] == "excluido_2025":
            verificar = r["senal"]
        elif r["tipo_baseline"].startswith("nuevo"):
            verificar = _hint(r["notas"]) or "confirmar elegibilidad con evidencia verbatim"
        else:
            verificar = "re-verificar actividad 2026 y recodificar a esquema 2026"
        enr.append((pr, disp, r, verificar))
    enr.sort(key=lambda x: (x[0], x[2]["pais"], x[2]["nombre_tentativo"]))

    ws = wb.create_sheet("Candidatos"); ws.sheet_properties.tabColor = "2E75B6"
    headers = ["Prioridad", "candidato_id", "País", "Nombre", "Tipo", "Canal",
               "Fuente", "URL(s)", "Señal / evidencia", "Qué verificar",
               "¿Elegible? (SI/NO/DUDA)", "Decisión", "Tipo de proceso (2026)",
               "Tipo convocante (2026)", "Nivel (0-3)", "Revisor", "Comentarios"]
    ws.append(headers)
    for pr, disp, r, verificar in enr:
        fuente = ""
        for tok in (r["notas"] or "").split(";"):
            if tok.strip().lower().startswith("fuente="):
                fuente = tok.split("=", 1)[1].strip()
        ws.append([pr, r["candidato_id"], r["pais"], r["nombre_tentativo"], disp,
                   r["canal_origen"], fuente, r["urls"], r["senal"], verificar,
                   "", "", "", "", "", "", ""])
    n = ws.max_row
    _estilo_header(ws, len(headers), dec_from=11)
    _widths(ws, [8, 30, 6, 40, 22, 12, 12, 40, 60, 45, 14, 16, 22, 22, 9, 14, 30])
    for col in (4, 8, 9, 10, 13, 14, 17):  # wrap en columnas largas
        for row in range(2, n + 1):
            ws.cell(row=row, column=col).alignment = WRAP
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n}"

    dv1 = DataValidation(type="list", formula1='"SI,NO,DUDA"', allow_blank=True)
    dv2 = DataValidation(type="list", formula1='"Aprobar Fase 2,Verificar,Descartar,Mantener"', allow_blank=True)
    dv3 = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    for dv, col in ((dv1, "K"), (dv2, "L"), (dv3, "O")):
        ws.add_data_validation(dv); dv.add(f"{col}2:{col}{n}")
    return len(enr)


def hoja_recodificacion(wb):
    rows = _read("data/baseline/recodificacion_2025.csv")
    ws = wb.create_sheet("Recodificación 2025"); ws.sheet_properties.tabColor = "BF9000"
    headers = ["País", "Caso", "Convocante (borrador)", "Nivel (borrador)",
               "rol_IA (borrador)", "Nota / fuente",
               "Convocante CONFIRMADO", "Nivel CONFIRMADO (0-3)",
               "¿Subir a 3? (evidencia recurrencia)", "Revisor", "Comentario"]
    ws.append(headers)
    for r in rows:
        nota = "; ".join(x for x in [r.get("fuente_convocante", ""), r.get("nota", "")] if x)
        ws.append([r["pais"], r["caso"], r.get("tipo_convocante", ""),
                   r.get("nivel_consolidacion", ""), r.get("rol_IA", ""), nota,
                   "", "", "", "", ""])
    n = ws.max_row
    _estilo_header(ws, len(headers), dec_from=7)
    _widths(ws, [6, 46, 30, 9, 12, 50, 30, 12, 22, 14, 30])
    for col in (2, 3, 6, 7):
        for row in range(2, n + 1):
            ws.cell(row=row, column=col).alignment = WRAP
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n}"
    dv = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f"H2:H{n}")


def hoja_medios(wb):
    rows = _read("data/baseline/directorio_medios.csv")
    if not rows:
        return
    ws = wb.create_sheet("Medios (A.10)"); ws.sheet_properties.tabColor = "548235"
    headers = ["País", "Medio", "Tipo", "URL", "Relevancia"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get("pais", ""), r.get("medio", ""), r.get("tipo", ""),
                   r.get("url", ""), r.get("relevancia", "")])
    n = ws.max_row
    _estilo_header(ws, len(headers))
    _widths(ws, [6, 28, 22, 40, 70])
    for row in range(2, n + 1):
        ws.cell(row=row, column=5).alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{n}"


def hoja_resumen(wb):
    rows = _read("data/candidatos/candidatos.csv")
    paises = ["AR","BO","BR","CL","CO","CR","CU","DO","EC","GT","HN","JM","MX","PA","PE","PY","SV","TT","UY","VE"]
    def cnt(p, pred): return sum(1 for r in rows if r["pais"] == p and pred(r))
    ws = wb.create_sheet("Resumen"); ws.sheet_properties.tabColor = "7030A0"
    headers = ["País", "Baseline activo", "Reverificados 2026", "NUEVOS", "NUEVOS dudosos", "Excluidos 2025"]
    ws.append(headers)
    for p in paises:
        ws.append([
            p,
            cnt(p, lambda r: r["tipo_baseline"] == "activo_2025"),
            cnt(p, lambda r: r["tipo_baseline"] == "activo_2025" and "reconfirmado web" in (r["notas"] or "")),
            cnt(p, lambda r: r["tipo_baseline"] == "nuevo"),
            cnt(p, lambda r: r["tipo_baseline"] == "nuevo_dudoso"),
            cnt(p, lambda r: r["tipo_baseline"] == "excluido_2025"),
        ])
    tot = ["TOTAL"] + [sum(ws.cell(row=r, column=c).value for r in range(2, ws.max_row + 1)) for c in range(2, 7)]
    ws.append(tot)
    n = ws.max_row
    _estilo_header(ws, len(headers))
    _widths(ws, [8, 16, 18, 10, 16, 16])
    for c in range(1, 7):
        ws.cell(row=n, column=c).font = Font(bold=True)
    ws.freeze_panes = "B2"


def main(argv=None):
    fecha = datetime.date.today().isoformat()
    wb = Workbook()
    hoja_instrucciones(wb, fecha)
    n = hoja_candidatos(wb)
    hoja_recodificacion(wb)
    hoja_medios(wb)
    hoja_resumen(wb)
    dst = os.path.join(ROOT, "data", "gates", "validacion_candidatos_ILIA2026.xlsx")
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    wb.save(dst)
    print(f"[ok] {dst}  ({n} candidatos + recodificación + medios + resumen)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
