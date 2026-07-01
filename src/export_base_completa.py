#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_base_completa.py — Base completa ILIA 2026 con las validaciones aplicadas.

Tres pestañas limpias:
  1 · Casos que entran   -> todo el detalle por caso + «¿Qué validar y cómo?» por caso
  2 · Excluidos          -> país, caso, motivo (lenguaje normal), fuente
  3 · Evidencia y fuentes-> cita textual por campo crítico + URLs (anti-alucinación)

Salida: data/gates/base_completa_ILIA2026.xlsx
"""
from __future__ import annotations
import glob, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import export_planilla_simple as eps       # clasificación (EXCLUIR, is_in, is_duda)
import export_por_pais as ep                # limpiar() para motivos en lenguaje normal

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
AZUL=PatternFill("solid",fgColor="1F4E78"); VERDE=PatternFill("solid",fgColor="C6EFCE")
AMBARH=PatternFill("solid",fgColor="FFF2CC"); GRIS=PatternFill("solid",fgColor="D9D9D9")
WB=Font(bold=True,color="FFFFFF"); CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
TL=Alignment(horizontal="left",vertical="top",wrap_text=True)
TH=Side(style="thin",color="BFBFBF"); BOR=Border(left=TH,right=TH,top=TH,bottom=TH)

CRIT=["pais","tipo_proceso","etapas_uso_IA","estado_actividad","nivel_consolidacion","tipo_convocante","usa_IA_en_proceso","rol_IA"]
CRIT_LBL={"pais":"País","tipo_proceso":"Tipo de proceso","etapas_uso_IA":"Etapas uso IA",
 "estado_actividad":"Estado actividad","nivel_consolidacion":"Nivel (0-3)","tipo_convocante":"Convocante",
 "usa_IA_en_proceso":"Usa IA en el proceso","rol_IA":"Rol IA"}

# Instrucciones específicas para casos sensibles (además de las derivadas de los gaps).
SPECIAL={
 "TT-engagett-plataforma-go-vocal":"Confirmar si el módulo de IA de Go Vocal (Sensemaking) se activó sobre los aportes, o si el uso de IA es solo el chatbot informativo AskNDTS. Revisar el panel de engage.gov.tt o consultar al Ministerio (MPAAI).",
 "CL-jornada-de-escucha-lanzamiento-instituto":"Adjuntar el enlace directo del informe/dashboard de la jornada; el de Tableau no es verificable públicamente. El equipo ejecutó el proceso, confirmar la técnica de análisis usada.",
 "CU-consulta-popular-codigo-de-las-familias":"Confirmar la técnica exacta de agrupamiento de GEMA-CEN (similitud/clustering) para precisar tipos de IA.",
 "CO-gaitana-ia-resguardo-zenu":"Confirmar el alcance del uso del LLM (DeepSeek) sobre los aportes y el estado de consolidación (¿puntual o recurrente?).",
}


def j(v):
    if v is None: return ""
    if isinstance(v,list): return "; ".join(str(x) for x in v if x not in (None,""))
    return str(v)

def estado(d):
    cid=d["caso_id"]
    if cid in eps.EXCLUIR or not eps.is_in(d): return "NO"
    if cid in eps.RESCATE_ENTRA: return "SÍ"
    if cid in eps.RECLASIF_DUDA or cid in eps.DUDA_REVISAR: return "DUDA"
    if eps.is_duda(d): return "DUDA"
    return "SÍ"

def fuente(d):
    for s in (d.get("fuentes") or []):
        u=(s.get("url") or "").strip()
        if u.startswith("http"): return u
    return ""

FIELD_LBL={
 "nivel_consolidacion":"nivel de consolidación","usa_IA_en_proceso":"uso de IA en el proceso",
 "tipos_IA_familia":"tipos de IA","usa_LLM_generativa":"uso de LLM generativa",
 "estado_actividad":"estado de actividad","rol_IA":"rol de la IA","etapas_uso_IA":"etapas de uso de IA",
 "tipo_convocante":"convocante","tipo_proceso":"tipo de proceso","tipo_desarrollador_IA":"desarrollador de IA",
 "evidencia_actividad_2026":"evidencia de actividad 2026","ano_inicio":"año de inicio",
 "evidencia_verbatim":"cita textual","nivel_consolidacion / usa_IA_en_proceso":"nivel de consolidación y uso de IA",
}
def _friendly(campo):
    for k,v in sorted(FIELD_LBL.items(),key=lambda kv:-len(kv[0])): campo=campo.replace(k,v)
    return campo

def que_validar(d):
    """Instrucción por caso: qué revisar (si hay dudas en algún dato) y cómo."""
    cid=d["caso_id"]; partes=[]
    if cid in SPECIAL: partes.append(SPECIAL[cid])
    for g in (d.get("gaps") or []):
        campo=_friendly((g.get("campo") or "").strip())
        url=(g.get("url_para_buscar") or "").strip()
        u=url if url.startswith("http") else ""
        if campo:
            partes.append(f"Confirmar {campo}"+(f" → {u}" if u else " (abrir la fuente)"))
    conf=(d.get("confianza") or "")
    if conf.upper().startswith("MEDIA"):
        partes.append("Confianza media: cotejar tipo de proceso, etapa, nivel (0-3) y convocante con la fuente antes de calcular.")
    if not partes:
        partes.append("Datos completos y con fuente. Verificación estándar: abrir la fuente y confirmar tipo de proceso, etapas, nivel (0-3) y convocante.")
    return " · ".join(partes)


def hdr(ws,n,h=30):
    for c in range(1,n+1):
        cell=ws.cell(row=1,column=c); cell.fill=AZUL; cell.font=WB; cell.alignment=CEN; cell.border=BOR
    ws.row_dimensions[1].height=h; ws.freeze_panes="C2" if n>3 else "A2"
    ws.auto_filter.ref=f"A1:{get_column_letter(n)}{ws.max_row}"

def widths(ws,wl):
    for i,w in enumerate(wl,1): ws.column_dimensions[get_column_letter(i)].width=w

def body(ws,ncol):
    for r in range(2,ws.max_row+1):
        for c in range(1,ncol+1):
            ws.cell(row=r,column=c).alignment=TL; ws.cell(row=r,column=c).border=BOR


def main():
    fichas=[json.load(open(f,encoding="utf-8")) for f in sorted(glob.glob(os.path.join(ROOT,"data/fichas/*_detalle.json")))]
    fichas.sort(key=lambda d:(str(d.get("pais") or ""),str(d.get("nombre_caso") or "")))
    IN=[d for d in fichas if estado(d)=="SÍ"]
    OUT=[d for d in fichas if estado(d)=="NO"]
    wb=Workbook(); wb.remove(wb.active)

    # ---------- 1 · Casos que entran ----------
    ws=wb.create_sheet("1 · Casos que entran")
    cols=["País","Caso","Proceso participativo (descripción)","Tipo de proceso","Etapas uso IA",
          "Nivel (0-3)","Continuidad","Convocante (7-cat)","Tipo organización","Rol IA","Función de la IA",
          "Usa IA en el proceso","Sistema(s) de IA","Tipos / familia de IA","Usa LLM generativa",
          "Tipo desarrollador IA","Origen desarrollador","Año inicio","Estado actividad","Organización a cargo",
          "Cuenta como iniciativa","Confianza","Nº gaps","¿Qué validar y cómo?","Fuente principal (URL)","caso_id"]
    ws.append(cols)
    for d in IN:
        ws.append([d.get("pais") or "", d.get("nombre_caso") or "",
            ep.limpiar(d.get("proceso_participativo") or ""), j(d.get("tipo_proceso")), j(d.get("etapas_uso_IA")),
            "" if d.get("nivel_consolidacion") is None else d.get("nivel_consolidacion"), j(d.get("continuidad")),
            j(d.get("tipo_convocante")), j(d.get("tipo_organizacion")), d.get("rol_IA") or "",
            ep.limpiar(d.get("funcion_IA") or ""), d.get("usa_IA_en_proceso") or "", d.get("sistema_IA") or "",
            j(d.get("tipos_IA_familia")), d.get("usa_LLM_generativa") or "", j(d.get("tipo_desarrollador_IA")),
            d.get("origen_desarrollador") or "", j(d.get("ano_inicio")), d.get("estado_actividad") or "",
            d.get("organizacion_a_cargo") or "", "NO (duplicado)" if d.get("cuenta_como_iniciativa") is False else "sí",
            d.get("confianza") or "", len(d.get("gaps") or []), que_validar(d), fuente(d), d.get("caso_id")])
    hdr(ws,len(cols))
    widths(ws,[6,30,50,22,20,7,14,22,14,12,42,14,26,22,12,18,14,9,16,26,14,26,6,60,44,30]); body(ws,len(cols))
    ci={h:i+1 for i,h in enumerate(cols)}
    for r in range(2,ws.max_row+1):
        g=ws.cell(row=r,column=ci["Nº gaps"]); g.alignment=CEN
        if isinstance(g.value,int) and g.value>0: g.fill=AMBARH
        ws.cell(row=r,column=ci["¿Qué validar y cómo?"]).fill=AMBARH

    # ---------- 2 · Excluidos ----------
    ws2=wb.create_sheet("2 · Excluidos")
    c2=["País","Caso","Motivo","Fuente (URL)","caso_id"]
    ws2.append(c2)
    for d in sorted(OUT,key=lambda x:(str(x.get("pais")),str(x.get("nombre_caso")))):
        cid=d["caso_id"]
        mot=eps.EXCLUIR.get(cid) or d.get("motivo_exclusion_2025") or d.get("justificacion_elegibilidad") or ""
        ws2.append([d.get("pais") or "", d.get("nombre_caso") or "", ep.limpiar(mot), fuente(d), cid])
    hdr(ws2,len(c2)); widths(ws2,[6,34,80,46,30]); body(ws2,len(c2))

    # ---------- 3 · Evidencia y fuentes ----------
    ws3=wb.create_sheet("3 · Evidencia y fuentes")
    c3=["País","Caso"]+[f"Evidencia: {CRIT_LBL[c]}" for c in CRIT]+["Fuentes (URLs)"]
    ws3.append(c3)
    for d in IN:
        ev=d.get("evidencia_verbatim",{}) or {}
        fu="\n".join(s.get("url","") for s in (d.get("fuentes") or []) if (s.get("url") or "").startswith("http"))
        ws3.append([d.get("pais") or "", d.get("nombre_caso") or ""]+[ev.get(c) or "" for c in CRIT]+[fu])
    hdr(ws3,len(c3)); widths(ws3,[6,30]+[30]*len(CRIT)+[46]); body(ws3,len(c3))
    for r in range(2,ws3.max_row+1):
        for c in range(3,3+len(CRIT)):
            if not ws3.cell(row=r,column=c).value: ws3.cell(row=r,column=c).fill=AMBARH

    dst=os.path.join(ROOT,"data/gates/base_completa_ILIA2026.xlsx"); wb.save(dst)
    dup=sum(1 for d in IN if d.get("cuenta_como_iniciativa") is False)
    print(f"[ok] {dst}")
    print(f"     entran: {len(IN)} (iniciativas efectivas: {len(IN)-dup}) · excluidos: {len(OUT)}")
    print(f"     casos con «qué validar» por gaps: {sum(1 for d in IN if d.get('gaps'))} · celdas de evidencia vacías se marcan en ámbar")
    return 0


if __name__ == "__main__":
    sys.exit(main())
