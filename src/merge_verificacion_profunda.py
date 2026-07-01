#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
merge_verificacion_profunda.py
================================
Integra los resultados de la VERIFICACIÓN PROFUNDA (fuentes primarias + cita
verbatim) de los 15 casos marcados, en:
  1. la planilla consolidada (columnas nuevas + hoja 'Verificación profunda'),
  2. la recodificación (rol_IA verificado, ya con cita verbatim -> cumple la
     regla del proyecto),
  3. recálculo del conjunto elegible.

Entrada: scratchpad/deep/*.json  (un objeto por caso, del verificador)
Salidas: data/gates/Planilla_Consolidada_Casos_ILIA2026.xlsx (in-place),
         data/gates/verificacion_profunda_resultados.json (copia en repo)
"""
import os
import re
import csv
import glob
import json
import unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCRATCH = ("/tmp/claude-0/-home-user-ilia26-participacion/"
           "f142e11e-c75e-5a60-95ab-939ca0282af8/scratchpad")
DEEP_DIR = os.path.join(SCRATCH, "deep")
XLSX = "data/gates/Planilla_Consolidada_Casos_ILIA2026.xlsx"
RECOD = "data/baseline/recodificacion_2025.csv"
OUT_JSON = "data/gates/verificacion_profunda_resultados.json"


def norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", s)).strip()


# --- cargar resultados profundos --------------------------------------------
results = []
for f in sorted(glob.glob(os.path.join(DEEP_DIR, "*.json"))):
    results.append(json.load(open(f, encoding="utf-8")))
print(f"Resultados profundos cargados: {len(results)}")

# decisión: rol_IA en {apoyo, sin_IA} => EXCLUIR ; contenido => INCLUIR ;
# no_concluyente => MANTENER_MARCADO
def decision(r):
    rol = (r.get("rol_IA_determinado") or "").lower()
    if rol == "contenido":
        return "INCLUIR"
    if rol in ("apoyo", "sin_ia", "sin ia"):
        return "EXCLUIR"
    return "MANTENER_MARCADO"

for r in results:
    r["_decision"] = decision(r)

by_key = {(r.get("pais", ""), norm(r.get("caso", ""))): r for r in results}


def lookup(pais, caso):
    nc = norm(caso)
    if (pais, nc) in by_key:
        return by_key[(pais, nc)]
    for (p, c), r in by_key.items():
        if p == pais and (nc[:20] in c or c[:20] in nc):
            return r
    return None


# --- 1) planilla: columnas nuevas en 'Detalle consolidado' ------------------
wb = openpyxl.load_workbook(XLSX)
wsd = wb["Detalle consolidado"]
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DEC_COLOR = {"INCLUIR": "548235", "EXCLUIR": "C00000", "MANTENER_MARCADO": "BF8F00"}
ROL_COLOR = {"contenido": "548235", "apoyo": "C00000", "sin_IA": "C00000", "no_concluyente": "BF8F00"}

start = wsd.max_column + 1
new_cols = [("rol_IA verificado (prof.)", 16), ("Recomendación final", 18),
            ("Confianza", 10), ("Cita verbatim (fuente primaria)", 62),
            ("Fuente de la cita", 40)]
for off, (h, w) in enumerate(new_cols):
    c = wsd.cell(row=1, column=start + off, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER; c.alignment = CENTER
    wsd.column_dimensions[get_column_letter(start + off)].width = w

excluir, incluir, marcar = [], [], []
for i in range(2, wsd.max_row + 1):
    pais = wsd.cell(row=i, column=2).value
    caso = wsd.cell(row=i, column=3).value
    r = lookup(pais, caso)
    if not r:
        continue
    rol = r.get("rol_IA_determinado", "")
    dec = r["_decision"]
    vals = [rol, dec, r.get("confianza", ""), r.get("cita_verbatim", ""), r.get("cita_fuente_url", "")]
    for off, v in enumerate(vals):
        c = wsd.cell(row=i, column=start + off, value=v)
        c.border = BORDER
        c.alignment = CENTER if off in (0, 1, 2) else WRAP
    if rol in ROL_COLOR:
        wsd.cell(row=i, column=start).font = Font(bold=True, color=ROL_COLOR[rol])
    if dec in DEC_COLOR:
        wsd.cell(row=i, column=start + 1).font = Font(bold=True, color=DEC_COLOR[dec])
    tup = (pais, caso, rol, r.get("confianza", ""))
    (excluir if dec == "EXCLUIR" else incluir if dec == "INCLUIR" else marcar).append(tup)

# --- 2) hoja 'Verificación profunda' ----------------------------------------
if "Verificación profunda" in wb.sheetnames:
    del wb["Verificación profunda"]
wv = wb.create_sheet("Verificación profunda")
wv["A1"] = "Verificación profunda (fuentes primarias + cita verbatim)"
wv["A1"].font = Font(bold=True, size=14, color="1F4E78")
wv["A2"] = ("Determinación del rol_IA de los 15 casos marcados. rol_IA ∈ {apoyo, sin_IA} ⇒ "
            "EXCLUIR (la IA no procesa el contenido participativo); contenido ⇒ INCLUIR.")
wv["A2"].font = Font(italic=True, size=10, color="595959")
cols = [("País", 6), ("Caso", 46), ("rol_IA verif.", 14), ("Recomendación", 16),
        ("Confianza", 10), ("Cita verbatim", 70), ("Fuente", 40)]
for j, (h, w) in enumerate(cols, 1):
    c = wv.cell(row=4, column=j, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER; c.alignment = CENTER
    wv.column_dimensions[get_column_letter(j)].width = w
row = 5
for r in sorted(results, key=lambda z: (z["_decision"], z.get("pais", ""))):
    vals = [r.get("pais", ""), r.get("caso", ""), r.get("rol_IA_determinado", ""),
            r["_decision"], r.get("confianza", ""), r.get("cita_verbatim", ""),
            r.get("cita_fuente_url", "")]
    for j, v in enumerate(vals, 1):
        c = wv.cell(row=row, column=j, value=v)
        c.border = BORDER
        c.alignment = CENTER if j in (1, 3, 4, 5) else WRAP
    rol = r.get("rol_IA_determinado", "")
    if rol in ROL_COLOR:
        wv.cell(row=row, column=3).font = Font(bold=True, color=ROL_COLOR[rol])
    if r["_decision"] in DEC_COLOR:
        wv.cell(row=row, column=4).font = Font(bold=True, color=DEC_COLOR[r["_decision"]])
    row += 1
wv.freeze_panes = "A5"

# recuento y conjunto elegible recomendado
n_base_in = sum(1 for i in range(2, wsd.max_row + 1)
                if str(wsd.cell(row=i, column=1).value).startswith("1."))
n_new_in = sum(1 for i in range(2, wsd.max_row + 1)
               if str(wsd.cell(row=i, column=1).value).startswith("5."))
total_A_prev = n_base_in + n_new_in
excl_A = [t for t in excluir]  # todos los excluir recomendados
# de los excluidos, cuántos contaban bajo A (todos los marcados estaban dentro)
nuevo_total_A = total_A_prev - len(excl_A)
row += 1
wv.cell(row=row, column=1, value="Resumen").font = Font(bold=True, size=12, color="1F4E78")
row += 1
for txt in [
    f"Casos verificados: {len(results)}",
    f"Recomendados EXCLUIR (rol_IA=apoyo/sin_IA): {len(excluir)}",
    f"Recomendados INCLUIR (rol_IA=contenido): {len(incluir)}",
    f"MANTENER_MARCADO (no concluyente): {len(marcar)}",
    f"Conjunto elegible A antes: {total_A_prev}  ->  aplicando recomendación: {nuevo_total_A}",
    f"(Bajo Escenario B, restar además DO CiudadanIA y CL TQHCH si siguen dentro.)",
]:
    wv.cell(row=row, column=1, value=txt).font = Font(size=10)
    row += 1

wb.save(XLSX)

# --- 3) recodificación: fijar rol_IA verificado (con respaldo verbatim) ------
# La verificación profunda SUPERSEDE las anotaciones "VALIDACIÓN 2026" previas.
rows = list(csv.reader(open(RECOD, encoding="utf-8")))
hdr = rows[0]
pi, ci, ri, ni = hdr.index("pais"), hdr.index("caso"), hdr.index("rol_IA"), hdr.index("nota")
changed = 0


def strip_validacion(nota):
    # elimina los segmentos "| VALIDACIÓN 2026: ..." (superados por la verif. profunda)
    return re.sub(r"\s*\|\s*VALIDACIÓN 2026:[^|]*", "", nota).strip()


for rrow in rows[1:]:
    if len(rrow) <= ni:
        continue
    r = lookup(rrow[pi], rrow[ci])
    if not r:
        continue
    rol = (r.get("rol_IA_determinado") or "").lower()
    conf = r.get("confianza", "")
    cita = (r.get("cita_verbatim", "") or "")[:220]
    url = r.get("cita_fuente_url", "")
    if "VERIF. PROFUNDA 2026" in rrow[ni]:
        continue
    rrow[ni] = strip_validacion(rrow[ni])
    if rol in ("apoyo", "sin_ia", "sin ia"):
        new_rol = "apoyo" if rol == "apoyo" else "sin_IA"
        if rrow[ri] != new_rol:
            rrow[ri] = new_rol
            changed += 1
        marca = f"VERIF. PROFUNDA 2026 (conf. {conf}): rol_IA={new_rol} -> EXCLUIR. Verbatim: «{cita}» [{url}]"
    elif rol == "contenido":
        marca = f"VERIF. PROFUNDA 2026 (conf. {conf}): rol_IA=contenido CONFIRMADO -> INCLUIR. Verbatim: «{cita}» [{url}]"
    else:
        marca = f"VERIF. PROFUNDA 2026 (conf. {conf}): {r.get('rol_IA_determinado','')} -> MANTENER_MARCADO."
    rrow[ni] = (rrow[ni] + (" | " if rrow[ni] else "") + marca).strip()
with open(RECOD, "w", encoding="utf-8", newline="") as f:
    csv.writer(f).writerows(rows)

json.dump(results, open(OUT_JSON, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

print(f"Planilla actualizada. Recod rol_IA cambiados: {changed}")
print(f"EXCLUIR={len(excluir)} INCLUIR={len(incluir)} MARCAR={len(marcar)}")
print(f"Conjunto elegible A: {total_A_prev} -> {nuevo_total_A}")
print("\nEXCLUIR:")
for t in excluir:
    print(f"  [{t[0]}] {t[1][:44]:44} rol={t[2]} ({t[3]})")
if incluir:
    print("INCLUIR (contenido confirmado):")
    for t in incluir:
        print(f"  [{t[0]}] {t[1][:44]:44} rol={t[2]} ({t[3]})")
if marcar:
    print("MANTENER_MARCADO:")
    for t in marcar:
        print(f"  [{t[0]}] {t[1][:44]:44} rol={t[2]} ({t[3]})")
