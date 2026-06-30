#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
merge_validacion.py
====================
Fusiona los resultados de la validación web (realidad del caso + URL funcional,
detección de alucinaciones) en la planilla consolidada, añadiendo columnas de
validación a la hoja 'Detalle consolidado' y una hoja 'Validación' con el
resumen y las alertas.

Entrada: scratchpad/validacion_resultados.json  (array de objetos por caso)
Salida : data/gates/Planilla_Consolidada_Casos_ILIA2026.xlsx  (in-place)
"""
import os
import re
import json
import unicodedata
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCRATCH = ("/tmp/claude-0/-home-user-ilia26-participacion/"
           "f142e11e-c75e-5a60-95ab-939ca0282af8/scratchpad")
RES_JSON = os.path.join(SCRATCH, "validacion_resultados.json")
XLSX = "data/gates/Planilla_Consolidada_Casos_ILIA2026.xlsx"


def norm(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


res = json.load(open(RES_JSON, encoding="utf-8"))
by_key = {}
for r in res:
    by_key[(r.get("pais", ""), norm(r.get("caso", "")))] = r


def lookup(pais, caso):
    nc = norm(caso)
    if (pais, nc) in by_key:
        return by_key[(pais, nc)]
    for (p, c), r in by_key.items():
        if p == pais and (nc[:22] in c or c[:22] in nc):
            return r
    return None


wb = openpyxl.load_workbook(XLSX)
wsd = wb["Detalle consolidado"]

# estilos
HDR_FILL = PatternFill("solid", fgColor="1F4E78")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
VERD_COLOR = {"REAL_CONFIRMADO": "548235", "PARCIAL": "BF8F00",
              "NO_VERIFICABLE": "C00000", "CONTRADICHO": "C00000"}

# nuevas columnas al final de Detalle consolidado
start = wsd.max_column + 1
new_cols = [("Validación realidad", 18), ("Usa IA en participación", 16),
            ("URL funcional (verificada)", 40), ("Estado URL original", 18),
            ("Evidencia verificación", 56), ("Alerta validación", 40)]
for off, (h, w) in enumerate(new_cols):
    j = start + off
    c = wsd.cell(row=1, column=j, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER; c.alignment = CENTER
    wsd.column_dimensions[get_column_letter(j)].width = w

stats = {"REAL_CONFIRMADO": 0, "PARCIAL": 0, "NO_VERIFICABLE": 0, "CONTRADICHO": 0, "sin_dato": 0}
alertas = []
url_rotas = []
for i in range(2, wsd.max_row + 1):
    cat = str(wsd.cell(row=i, column=1).value or "")
    pais = wsd.cell(row=i, column=2).value
    caso = wsd.cell(row=i, column=3).value
    if not (cat.startswith("1.") or cat.startswith("5.")):
        continue  # solo casos incluidos
    r = lookup(pais, caso)
    if not r:
        stats["sin_dato"] += 1
        wsd.cell(row=i, column=start, value="(sin dato)").alignment = CENTER
        continue
    verd = r.get("veredicto_realidad", "")
    stats[verd] = stats.get(verd, 0) + 1
    vals = [verd, r.get("usa_IA_en_participacion", ""), r.get("url_funcional", ""),
            r.get("url_original_estado", ""), r.get("evidencia", ""), r.get("alerta", "")]
    for off, v in enumerate(vals):
        c = wsd.cell(row=i, column=start + off, value=v)
        c.border = BORDER
        c.alignment = CENTER if off in (0, 1, 3) else WRAP
    col = VERD_COLOR.get(verd)
    if col:
        wsd.cell(row=i, column=start).font = Font(bold=True, color=col)
    if r.get("alerta"):
        wsd.cell(row=i, column=start + 5).fill = PatternFill("solid", fgColor="FFE699")
        alertas.append((pais, caso, verd, r.get("alerta")))
    est = str(r.get("url_original_estado", ""))
    if "404" in est or "muerta" in est or not r.get("url_funcional"):
        url_rotas.append((pais, caso, est, r.get("url_funcional", "")))

# hoja Validación
if "Validación" in wb.sheetnames:
    del wb["Validación"]
wsv = wb.create_sheet("Validación")
wsv["A1"] = "Validación de casos incluidos — realidad, URL funcional y alertas"
wsv["A1"].font = Font(bold=True, size=14, color="1F4E78")
wsv["A2"] = ("Verificación web (WebSearch/WebFetch) de los 37 casos elegibles 2026. "
             "Objetivo: descartar alucinaciones y confirmar una URL operativa por caso.")
wsv["A2"].font = Font(italic=True, size=10, color="595959")
row = 4
wsv.cell(row=row, column=1, value="Resultado").font = Font(bold=True, size=12)
row += 1
for j, h in enumerate(["Veredicto realidad", "N° casos"], 1):
    c = wsv.cell(row=row, column=j, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER; c.alignment = CENTER
row += 1
for k in ["REAL_CONFIRMADO", "PARCIAL", "NO_VERIFICABLE", "CONTRADICHO", "sin_dato"]:
    wsv.cell(row=row, column=1, value=k).border = BORDER
    b = wsv.cell(row=row, column=2, value=stats.get(k, 0))
    b.border = BORDER; b.alignment = CENTER
    if k in VERD_COLOR:
        wsv.cell(row=row, column=1).font = Font(bold=True, color=VERD_COLOR[k])
    row += 1

row += 1
wsv.cell(row=row, column=1, value=f"Alertas de validación ({len(alertas)})").font = Font(bold=True, size=12, color="C00000")
row += 1
for j, (h, w) in enumerate([("País", 6), ("Caso", 46), ("Veredicto", 18), ("Alerta", 80)], 1):
    c = wsv.cell(row=row, column=j, value=h)
    c.fill = HDR_FILL; c.font = HDR_FONT; c.border = BORDER; c.alignment = CENTER
    wsv.column_dimensions[get_column_letter(j)].width = w
row += 1
for pais, caso, verd, al in alertas:
    wsv.cell(row=row, column=1, value=pais).border = BORDER
    wsv.cell(row=row, column=2, value=caso).border = BORDER
    wsv.cell(row=row, column=2).alignment = WRAP
    wsv.cell(row=row, column=3, value=verd).border = BORDER
    wsv.cell(row=row, column=4, value=al).border = BORDER
    wsv.cell(row=row, column=4).alignment = WRAP
    row += 1

row += 1
wsv.cell(row=row, column=1, value=f"URLs a corregir / sin URL funcional ({len(url_rotas)})").font = Font(bold=True, size=12, color="BF8F00")
row += 1
for pais, caso, est, urlf in url_rotas:
    wsv.cell(row=row, column=1, value=pais).border = BORDER
    wsv.cell(row=row, column=2, value=caso).border = BORDER
    wsv.cell(row=row, column=2).alignment = WRAP
    wsv.cell(row=row, column=3, value=est).border = BORDER
    wsv.cell(row=row, column=4, value=("Sugerida: " + urlf) if urlf else "Sin URL funcional encontrada").border = BORDER
    wsv.cell(row=row, column=4).alignment = WRAP
    row += 1
wsv.freeze_panes = "A4"

wb.save(XLSX)
print("OK merge ->", XLSX)
print("Stats:", stats)
print(f"Alertas: {len(alertas)} | URLs a corregir: {len(url_rotas)}")
for a in alertas:
    print("  ALERTA:", a[0], a[1][:40], "->", a[3][:80])
