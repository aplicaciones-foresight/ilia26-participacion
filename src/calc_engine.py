#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
calc_engine.py — Motor de cálculo ILIA · Participación Ciudadana.

Dos modos, ambos determinísticos y reproducibles (Regla de oro: ningún número
sale de memoria; todo se calcula desde la BBDD):

  · LEGACY 2025  -> compute_full()  : fórmula de 4 variables, validada 19/19 vs.
                    el oficial 2025 (≤1 punto). NO se modifica (§8).
  · 2026         -> compute_2026()  : 5 variables (§4.2), continuidad por nivel
                    máximo, convocante a dos niveles, cantidad por umbrales
                    parametrizados, escenarios A/B. Sub2 reusa el código 2025.

El cálculo SIEMPRE corre primero el test de no-regresión legacy (ver tests/).
"""
from __future__ import annotations
import os
import openpyxl

PAISES = ["AR","BO","BR","CL","CO","CR","CU","DO","EC","GT","HN","JM","MX","PA","PE","PY","SV","UY","VE"]

_THIS = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.normpath(os.path.join(_THIS, "..", "data", "baseline", "BBDD_IA_Participacion_2025.xlsx"))


# ============================================================================
# LOADER (path configurable; comportamiento de extracción de campos = 2025)
# ============================================================================
def load_casos(path: str | None = None, sheet: str = "BBDD Casos"):
    """Carga los casos de la BBDD 2025. Solo se parametriza la RUTA; el mapeo de
    columnas es idéntico al motor 2025 validado."""
    wb = openpyxl.load_workbook(path or DEFAULT_XLSX, data_only=True)
    ws = wb[sheet]
    casos = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None or row[1] is None:
            continue
        if str(row[0]).strip() == "" or str(row[1]).strip() == "":
            continue
        casos.append({
            "pais": str(row[0]).strip(), "caso": str(row[1]).strip(),
            "ano": row[2], "activo": str(row[3]).strip() if row[3] else "",
            "tipo_org": str(row[6]).strip() if row[6] else "",
            "tipos_ia": str(row[9]).strip() if row[9] else "",
            "tipo_proceso": str(row[11]).strip() if row[11] else "",
            "etapa": str(row[12]).strip() if row[12] else "",
            "desarrollador": str(row[14]).strip() if row[14] else "",
            "origen": str(row[15]).strip() if row[15] else "",
            "continuidad": str(row[17]).strip() if row[17] else "",
        })
    return casos


def split_clean(s):
    if isinstance(s, (list, tuple)):
        return [str(x).strip() for x in s if str(x).strip()]
    return [x.strip() for x in str(s).replace(";", ",").replace("|", ",").split(",") if x.strip()]


# ---- Normalizadores 2025 (legacy; NO tocar la semántica) -------------------
PROC_MAP = {"participación digital":"Participación digital","participacion digital":"Participación digital",
    "governanza colaborativa":"Governanza","governanza participativa":"Governanza",
    "mini públicos":"Mini públicos","mini publicos":"Mini públicos",
    "referendos e iniciativas populares":"Referendos","referendos":"Referendos",
    "presupuestos participativos":"Presupuestos participativos"}
def norm_proc(p): return PROC_MAP.get(p.lower().strip())
def norm_etapa(e):
    e = e.lower()
    if "planif" in e: return "Planificación"
    if "implement" in e: return "Implementación"
    if "análisis" in e or "analisis" in e: return "Análisis"
    if "traducción" in e or "traduccion" in e: return "Traducción Política"
    return None
def norm_org(o):
    o = o.lower()
    if "mixto" in o: return "Mixto"
    if "público" in o or "publico" in o: return "Público"
    if "privado" in o: return "Privado"
    return None
DEV_MAP = {"industria":"Empresa","academia":"Universidad","gobierno":"Gobierno","sociedad civil":"Sociedad Civil"}
def norm_dev(d): return DEV_MAP.get(d.lower().strip())


def raw_counts(casos_list):
    """Conteos crudos por país (compartido entre 2025 y Sub2 de 2026)."""
    out = {}
    for p in PAISES:
        cp = [c for c in casos_list if c["pais"] == p]
        tipos=set(); etapas=set(); cont=set(); orgs=set()
        dev_nac=set(); dev_int=set(); sistemas=set()
        for c in cp:
            for x in split_clean(c["tipo_proceso"]):
                m = norm_proc(x); tipos.add(m) if m else None
            for x in split_clean(c["etapa"]):
                m = norm_etapa(x); etapas.add(m) if m else None
            cc = str(c["continuidad"]).lower()
            if "único" in cc or "unico" in cc: cont.add("UU")
            if "plataforma" in cc: cont.add("PL")
            m = norm_org(c["tipo_org"]); orgs.add(m) if m else None
            devs = [norm_dev(x) for x in split_clean(c["desarrollador"])]
            devs = [d for d in devs if d]
            origs = [o.lower() for o in split_clean(c["origen"])]
            has_int = any("internacional" in o for o in origs)  # noqa: F841
            has_nac = any(o == "nacional" for o in origs)
            has_int = any(o == "internacional" for o in origs)
            for d in devs:
                if has_nac: dev_nac.add(d)
                if has_int: dev_int.add(d)
            for s in split_clean(c["tipos_ia"]):
                sistemas.add(s.lower())
        out[p] = {"tipos":len(tipos),"etapas":len(etapas),"cont":len(cont),"orgs":len(orgs),
                  "dev_nac":len(dev_nac),"dev_int":len(dev_int),"sistemas":len(sistemas),"n":len(cp)}
    return out


# ============================================================================
# MODO LEGACY 2025 — fórmula de 4 variables (INTACTA, §8)
# ============================================================================
def compute_full(casos_list, estado_func=None, use_5var=False):
    """Calcula sub1, sub2, indicador 2025. estado_func(caso)->peso si use_5var."""
    rc = raw_counts(casos_list)
    sub1 = {}
    for p in PAISES:
        d = rc[p]
        if d["n"] == 0:
            sub1[p] = {"v":0,"tipos":0,"etapas":0,"cont":0,"pp":0,"estado":0}; continue
        v_tipos=d["tipos"]/5*100; v_etapas=d["etapas"]/4*100
        v_cont=d["cont"]/2*100; v_pp=d["orgs"]/3*100
        if use_5var:
            cp=[c for c in casos_list if c["pais"]==p]
            pesos=[estado_func(c) for c in cp]
            v_estado=sum(pesos)/len(pesos)*100
            v=(v_tipos+v_etapas+v_cont+v_pp+v_estado)/5
            sub1[p]={"v":round(v),"tipos":round(v_tipos),"etapas":round(v_etapas),
                     "cont":round(v_cont),"pp":round(v_pp),"estado":round(v_estado)}
        else:
            v=(v_tipos+v_etapas+v_cont+v_pp)/4
            sub1[p]={"v":round(v),"tipos":round(v_tipos),"etapas":round(v_etapas),
                     "cont":round(v_cont),"pp":round(v_pp),"estado":None}
    max_nac=max((rc[p]["dev_nac"] for p in PAISES), default=0) or 1
    max_int=max((rc[p]["dev_int"] for p in PAISES), default=0) or 1
    max_sis=max((rc[p]["sistemas"] for p in PAISES), default=0) or 1
    sub2={}
    for p in PAISES:
        d=rc[p]
        if d["n"]==0:
            sub2[p]={"v":0,"dev":0,"sis":0}; continue
        nac=d["dev_nac"]/max_nac*100
        intl=d["dev_int"]/max_int*100
        dev=(nac+intl)/2
        sis=d["sistemas"]/max_sis*100
        v=(dev+sis)/2
        sub2[p]={"v":round(v),"dev":round(dev),"sis":round(sis)}
    ind={}
    for p in PAISES:
        ind[p]=round((sub1[p]["v"]+sub2[p]["v"])/2)
    return sub1,sub2,ind,rc


# ============================================================================
# TAXONOMÍAS 2026 (§4.2) y normalizadores tolerantes (§8)
# ============================================================================
PROC_MAP_2026 = {
    "participación digital":"Participación Digital","participacion digital":"Participación Digital",
    "gobernanza colaborativa":"Gobernanza Colaborativa","governanza colaborativa":"Gobernanza Colaborativa",
    "gobernanza participativa":"Gobernanza Colaborativa","governanza participativa":"Gobernanza Colaborativa",
    "governanza":"Gobernanza Colaborativa","gobernanza":"Gobernanza Colaborativa",
    "mini públicos":"Mini públicos","mini publicos":"Mini públicos","minipublicos":"Mini públicos",
    "referendos e iniciativas populares":"Referendos e Iniciativas Populares",
    "referendos":"Referendos e Iniciativas Populares","iniciativas populares":"Referendos e Iniciativas Populares",
    "presupuestos participativos":"Presupuestos Participativos","presupuesto participativo":"Presupuestos Participativos",
}
TIPOS_PROCESO_2026 = ["Participación Digital","Gobernanza Colaborativa","Mini públicos",
                      "Referendos e Iniciativas Populares","Presupuestos Participativos"]
def norm_proc_2026(p):
    return PROC_MAP_2026.get(str(p).lower().strip())

CONV_GUB = ["gobierno local","gobierno nacional","otra institución pública"]
CONV_NOGUB = ["empresa","universidad","sociedad civil","organización internacional"]
CONV_MAP = {
    "gobierno local":"gobierno local","gob local":"gobierno local","municipal":"gobierno local",
    "municipio":"gobierno local","gobierno municipal":"gobierno local","alcaldía":"gobierno local",
    "alcaldia":"gobierno local","prefeitura":"gobierno local","ayuntamiento":"gobierno local",
    "gobierno nacional":"gobierno nacional","gob nacional":"gobierno nacional","nacional":"gobierno nacional",
    "estado nacional":"gobierno nacional","gobierno federal":"gobierno nacional","federal":"gobierno nacional",
    "otra institución pública":"otra institución pública","otra institucion publica":"otra institución pública",
    "institución pública":"otra institución pública","institucion publica":"otra institución pública",
    "organismo público":"otra institución pública","organismo publico":"otra institución pública",
    "empresa":"empresa","industria":"empresa","privada":"empresa","sector privado":"empresa","govtech":"empresa",
    "universidad":"universidad","academia":"universidad","universidades":"universidad",
    "sociedad civil":"sociedad civil","ong":"sociedad civil","fundación":"sociedad civil","fundacion":"sociedad civil",
    "organización internacional":"organización internacional","organizacion internacional":"organización internacional",
    "org internacional":"organización internacional","org. internacional":"organización internacional",
    "organismo internacional":"organización internacional","multilateral":"organización internacional",
}
def norm_convocante(c):
    return CONV_MAP.get(str(c).lower().strip())


# ============================================================================
# RECODIFICACIÓN y ESCENARIOS (§6.3) para el modo 2026
# ============================================================================
def _key(pais, caso):
    return (str(pais).strip(), str(caso).strip().lower())


def aplicar_recodificacion(casos, recod_rows):
    """Fusiona la recodificación 2026 (tipo_convocante, nivel_consolidacion,
    rol_IA) sobre los casos 2025 por (país, caso). `recod_rows`: lista de dicts
    con claves pais, caso, tipo_convocante (str con '|' o ','), nivel_consolidacion,
    rol_IA. Casos sin recodificación quedan marcados con alerta."""
    idx = {_key(r["pais"], r["caso"]): r for r in recod_rows}
    out = []
    for c in casos:
        c = dict(c)
        r = idx.get(_key(c["pais"], c["caso"]))
        if r:
            c["tipo_convocante"] = split_clean(r.get("tipo_convocante", ""))
            niv = r.get("nivel_consolidacion", None)
            c["nivel_consolidacion"] = int(niv) if str(niv).strip() not in ("", "None", "nan") else None
            c["rol_IA"] = (r.get("rol_IA") or "").strip() or None
            c["_recodificado"] = True
        else:
            c["tipo_convocante"] = []
            c["nivel_consolidacion"] = None
            c["rol_IA"] = None
            c["_recodificado"] = False
        out.append(c)
    return out


def aplicar_escenario(casos, config):
    """Filtra casos según escenario activo: exclusiones firmes + adicionales del
    escenario, y descarta rol_IA == 'apoyo' (frontera §4.2)."""
    esc = (config or {}).get("escenarios", {})
    activo = esc.get("scenario_activo", "A")
    excl = list(esc.get("exclusiones_firmes", []) or [])
    excl += list((esc.get(activo, {}) or {}).get("exclusiones_adicionales", []) or [])
    excl_keys = {_key(e["pais"], e["caso"]) for e in excl}
    out = []
    for c in casos:
        if _key(c["pais"], c["caso"]) in excl_keys:
            continue
        if str(c.get("rol_IA", "")).strip().lower() == "apoyo":
            continue
        out.append(c)
    return out, {"scenario": activo, "n_excluidos_lista": len(excl_keys)}


def load_recodificacion_csv(path):
    """Lee la recodificación 2026 (BORRADOR pre-Gate). Resuelve el nivel efectivo
    usando 'nivel_consolidacion_override' si está presente. Devuelve filas listas
    para aplicar_recodificacion(). Si el archivo no existe, devuelve []."""
    import csv
    if not path or not os.path.exists(path):
        return []
    rows = []
    with open(path, newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            ov = (r.get("nivel_consolidacion_override") or "").strip()
            niv = ov if ov else (r.get("nivel_consolidacion") or "").strip()
            rows.append({
                "pais": r.get("pais", "").strip(),
                "caso": r.get("caso", "").strip(),
                "tipo_convocante": r.get("tipo_convocante", ""),
                "nivel_consolidacion": niv,
                "rol_IA": r.get("rol_IA", ""),
            })
    return rows


def run_2026(casos, config, recod_rows):
    """Pipeline 2026 completo y determinístico: recodifica -> aplica escenario ->
    calcula. Usado por simulate.py y reproduce.py para garantizar consistencia."""
    enriquecidos = aplicar_recodificacion(casos, recod_rows)
    filtrados, info = aplicar_escenario(enriquecidos, config)
    sub1, sub2, ind, rc = compute_2026(filtrados, config)
    info["n_casos_post_filtro"] = len(filtrados)
    return sub1, sub2, ind, rc, info


# ============================================================================
# MODO 2026 — 5 variables (§4.2)
# ============================================================================
def _as_list(x):
    if x is None: return []
    if isinstance(x, (list, tuple)): return [str(i).strip() for i in x if str(i).strip()]
    return split_clean(x)


def _valor_cantidad(n, thresholds):
    for t in thresholds:
        lo = t["min"]; hi = t["max"]
        if n >= lo and (hi is None or n <= hi):
            return t["value"]
    return 0


def _default_thresholds():
    return [{"min":0,"max":0,"value":0},{"min":1,"max":3,"value":25},
            {"min":4,"max":6,"value":50},{"min":7,"max":9,"value":75},
            {"min":10,"max":None,"value":100}]


def nivel_efectivo(caso, config):
    """Nivel de consolidación aplicando la regla de caducidad del nivel 1 (§4.2):
    un anuncio caduca a los N años sin implementación -> 0."""
    niv = caso.get("nivel_consolidacion")
    if niv is None:
        return None
    niv = int(niv)
    if niv == 1:
        cont = (config or {}).get("continuidad", {})
        cad = cont.get("caducidad_anuncio_anios", 2)
        ref = cont.get("anio_referencia_ciclo", 2026)
        ano = caso.get("ano")
        try:
            ano = int(float(ano)) if ano not in (None, "") else None
        except (ValueError, TypeError):
            ano = None
        if ano is not None and (ref - ano) > cad:
            return 0
    return niv


def compute_2026(casos, config):
    """Calcula Sub1 (5 var), Sub2 (idéntico a 2025) e indicador 2026 por país.

    `casos` deben venir ya recodificados y filtrados por escenario
    (ver aplicar_recodificacion + aplicar_escenario). Cada caso usa:
      Sub1: tipo_proceso, etapa, tipo_convocante, nivel_consolidacion, (ano)
      Sub2: desarrollador, origen, tipos_ia   (reusa raw_counts 2025)
    """
    cfg = config or {}
    cant = cfg.get("cantidad", {})
    thresholds = cant.get("thresholds") or _default_thresholds()
    count_min_level = cant.get("count_min_level", 0)
    if count_min_level in (None, ""):
        count_min_level = 0
    redondeo = (cfg.get("redondeo", {}) or {}).get("modo", "legacy")

    # ---- Sub1: 5 variables, promedio simple, máximo absoluto ----
    sub1 = {}
    for p in PAISES:
        cp = [c for c in casos if c["pais"] == p]
        if not cp:
            sub1[p] = {"v":0.0,"tipos":0.0,"etapas":0.0,"cont":0.0,"conv":0.0,"cant":0.0,"n":0}; continue

        tipos = set(); etapas = set()
        gub = set(); nogub = set(); niveles = []
        n_cant = 0
        for c in cp:
            for x in _as_list(c.get("tipo_proceso")):
                m = norm_proc_2026(x);  tipos.add(m) if m else None
            for x in _as_list(c.get("etapa", c.get("etapas_uso_IA"))):
                m = norm_etapa(x);      etapas.add(m) if m else None
            for x in _as_list(c.get("tipo_convocante")):
                m = norm_convocante(x)
                if m in CONV_GUB: gub.add(m)
                elif m in CONV_NOGUB: nogub.add(m)
            ne = nivel_efectivo(c, cfg)
            if ne is not None:
                niveles.append(ne)
                if ne >= count_min_level and ne > 0:
                    n_cant += 1

        v_tipos = len(tipos)/5*100
        v_etapas = len(etapas)/4*100
        v_cont = (max(niveles) if niveles else 0)/3*100
        v_conv = (len(gub)/3 + len(nogub)/4)/2*100
        v_cant = _valor_cantidad(n_cant, thresholds)
        v = (v_tipos + v_etapas + v_cont + v_conv + v_cant)/5
        sub1[p] = {"v":v,"tipos":v_tipos,"etapas":v_etapas,"cont":v_cont,
                   "conv":v_conv,"cant":v_cant,"n":len(cp),"n_cant":n_cant,
                   "nivel_max":(max(niveles) if niveles else 0)}

    # ---- Sub2: idéntico a 2025 (reusa raw_counts) ----
    rc = raw_counts(casos)
    max_nac = max((rc[p]["dev_nac"] for p in PAISES), default=0) or 1
    max_int = max((rc[p]["dev_int"] for p in PAISES), default=0) or 1
    max_sis = max((rc[p]["sistemas"] for p in PAISES), default=0) or 1
    sub2 = {}
    for p in PAISES:
        d = rc[p]
        if d["n"] == 0:
            sub2[p] = {"v":0.0,"dev":0.0,"sis":0.0}; continue
        nac = d["dev_nac"]/max_nac*100
        intl = d["dev_int"]/max_int*100
        dev = (nac+intl)/2
        sis = d["sistemas"]/max_sis*100
        sub2[p] = {"v":(dev+sis)/2,"dev":dev,"sis":sis}

    # ---- Indicador final con redondeo configurable (§4.4) ----
    ind = {}
    for p in PAISES:
        if redondeo == "legacy":
            s1 = round(sub1[p]["v"]); s2 = round(sub2[p]["v"])
            ind[p] = round((s1+s2)/2)
            sub1[p]["v_pub"] = s1; sub2[p]["v_pub"] = s2
        else:  # "solo_final"
            ind[p] = round((sub1[p]["v"]+sub2[p]["v"])/2)
            sub1[p]["v_pub"] = round(sub1[p]["v"]); sub2[p]["v_pub"] = round(sub2[p]["v"])
    return sub1, sub2, ind, rc


# ============================================================================
# CLI: corre el chequeo de no-regresión legacy (informativo).
# ============================================================================
if __name__ == "__main__":
    casos = load_casos()
    sub1, sub2, ind, rc = compute_full(casos)
    oficial_sub2={"AR":0,"BO":23,"BR":100,"CL":59,"CO":100,"CR":53,"CU":0,"DO":28,"EC":18,
                  "GT":33,"HN":52,"JM":0,"MX":37,"PA":0,"PE":51,"PY":0,"SV":0,"UY":0,"VE":0}
    oficial_ind={"AR":0,"BO":30,"BR":76,"CL":57,"CO":85,"CR":46,"CU":0,"DO":30,"EC":25,
                 "GT":32,"HN":42,"JM":0,"MX":47,"PA":0,"PE":53,"PY":0,"SV":0,"UY":0,"VE":0}
    print("VALIDACIÓN SUB2 e INDICADOR FINAL contra oficial 2025")
    print(f"{'País':<5}{'Sub2':>6}{'OfS2':>6}{'OK':>4}{'   ':<3}{'Ind':>5}{'OfInd':>7}{'OK':>4}")
    s2ok=True; iok=True
    for p in PAISES:
        o2=oficial_sub2[p]; oi=oficial_ind[p]
        c2=sub2[p]["v"]; ci=ind[p]
        k2=abs(c2-o2)<=1; ki=abs(ci-oi)<=1
        if not k2: s2ok=False
        if not ki: iok=False
        print(f"{p:<5}{c2:>6}{o2:>6}{'✓' if k2 else '✗':>4}{'   ':<3}{ci:>5}{oi:>7}{'✓' if ki else '✗':>4}")
    print(f"\nSub2: {'OK' if s2ok else 'DISCREPANCIAS'}  |  Indicador: {'OK' if iok else 'DISCREPANCIAS'}")
