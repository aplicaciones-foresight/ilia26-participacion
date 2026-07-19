#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_comparacion.py — Verificación del entregable Comparacion_2025_2026_ILIA.xlsx.

Recomputa en Python puro, desde el COMDING crudo de la pestaña «1 · BBDD 2025
(insumo)» (columnas de dato, no los flags), el cálculo 2025 con la fórmula
antigua (legacy 4 variables) y con la fórmula nueva (CENIA v2), y compara:

  · contra los valores cacheados de las pestañas 2 y 3 (tras recalc.py);
  · el indicador legacy contra el oficial 2025 publicado (±1; TT sin baseline);
  · la columna 2026 de la pestaña 4 contra el entregable
    BBDD_casos_incluidos_ILIA2026.xlsx (valores cacheados);
  · los máximos relativos del ciclo 2025 y la presencia de los 8 gráficos.

Sale con código ≠ 0 si algo no cuadra. Requiere haber corrido antes
scripts de recalc sobre ambos entregables.
"""
import os
import sys
import unicodedata
import zipfile
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

THIS = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.normpath(os.path.join(THIS, ".."))
OUT = os.path.join(ROOT, "data", "final", "Comparacion_2025_2026_ILIA.xlsx")
D1 = os.path.join(ROOT, "data", "final", "BBDD_casos_incluidos_ILIA2026.xlsx")
BBDD = os.path.join(ROOT, "data", "baseline", "BBDD_IA_Participacion_2025.xlsx")

PAISES = ["AR","BO","BR","CL","CO","CR","CU","DO","EC","GT","HN","JM","MX","PA","PE","PY","SV","TT","UY","VE"]
ANIO_REF, CADUCIDAD = 2026, 2

errors = []
def check(cond, msg):
    if not cond:
        errors.append(msg)
    return cond

def rhu(x):
    return int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))

def sin_acentos(s):
    return "".join(ch for ch in unicodedata.normalize("NFD", str(s).lower())
                   if not unicodedata.combining(ch))

def toks(s):
    return [t.strip().lower() for t in str(s or "").replace(";", ",").split(",") if t.strip()]

# --------------------------------------------------------------------------
# 1. Leer el coding crudo de la pestaña 1
# --------------------------------------------------------------------------
wb = openpyxl.load_workbook(OUT, data_only=True)
w1 = wb["1 · BBDD 2025 (insumo)"]
casos = []
for r in range(2, 30):
    if not w1.cell(r, 1).value:
        continue
    casos.append({
        "pais": str(w1.cell(r, 1).value).strip(),
        "ano": w1.cell(r, 3).value,
        "tipo_org": str(w1.cell(r, 6).value or ""),
        "proceso": toks(w1.cell(r, 7).value),
        "etapa": toks(w1.cell(r, 8).value),
        "cont": str(w1.cell(r, 9).value or ""),
        "dev": toks(w1.cell(r, 10).value),
        "origen": toks(w1.cell(r, 11).value),
        "sistemas": toks(w1.cell(r, 12).value),
        "conv": toks(w1.cell(r, 13).value),
        "nivel": w1.cell(r, 14).value,
    })
check(len(casos) == 28, f"pestaña 1: se esperaban 28 casos, hay {len(casos)}")

PROC5 = {"participacion digital": "PD", "governanza colaborativa": "GC",
         "gobernanza colaborativa": "GC", "mini publicos": "MP",
         "referendos e iniciativas populares": "RIP", "presupuestos participativos": "PP"}
GUB = ["gobierno local", "gobierno nacional", "otra institucion publica"]
NOGUB = ["empresa", "universidad", "sociedad civil", "organizacion internacional"]
DEV4 = ["industria", "academia", "gobierno", "sociedad civil"]

def etapa_id(t):
    t = sin_acentos(t)
    for sub, k in (("planif", "P"), ("implement", "I"), ("analisis", "A"), ("traduccion", "T")):
        if sub in t:
            return k
    return None

raw = {}
for p in PAISES:
    cp = [c for c in casos if c["pais"] == p]
    tipos = set(); etapas = set(); cont = set(); orgs = set()
    dnac = set(); dint = set(); sis = set()
    gub = set(); nogub = set(); combos = set(); niveles = []
    for c in cp:
        for x in c["proceso"]:
            m = PROC5.get(sin_acentos(x))
            check(m is not None, f"tipo de proceso no mapeado: {x!r}")
            if m: tipos.add(m)
        for x in c["etapa"]:
            m = etapa_id(x)
            check(m is not None, f"etapa no mapeada: {x!r}")
            if m: etapas.add(m)
        cc = sin_acentos(c["cont"])
        if "unico" in cc: cont.add("UU")
        if "plataforma" in cc: cont.add("PL")
        o = sin_acentos(c["tipo_org"])
        for sub, k in (("mixto", "M"), ("publico", "Pu"), ("privado", "Pr")):
            if sub in o: orgs.add(k); break
        has_nac = any(sin_acentos(x) == "nacional" for x in c["origen"])
        has_int = any(sin_acentos(x) == "internacional" for x in c["origen"])
        for d in c["dev"]:
            dn = sin_acentos(d)
            check(dn in DEV4, f"desarrollador no mapeado: {d!r}")
            if has_nac: dnac.add(dn)
            if has_int: dint.add(dn)
        sis.update(sin_acentos(x) for x in c["sistemas"])
        tt = set()
        for cv in c["conv"]:
            cvn = sin_acentos(cv)
            if cvn in GUB: gub.add(cvn); tt.add(cvn)
            elif cvn in NOGUB: nogub.add(cvn); tt.add(cvn)
            else: check(False, f"convocante no mapeado: {cv!r}")
        if len(tt) >= 2: combos.add(frozenset(tt))
        niv = c["nivel"]
        if niv is not None and str(niv).strip() != "":
            niv = int(float(niv))
            ano = c["ano"]
            if niv == 1 and ano not in (None, "") and (ANIO_REF - int(float(ano))) > CADUCIDAD:
                niv = 0
            niveles.append(niv)
    raw[p] = dict(n=len(cp), tipos=len(tipos), etapas=len(etapas), cont=len(cont),
                  orgs=len(orgs), dnac=len(dnac), dint=len(dint), sis=len(sis),
                  gub=len(gub), nogub=len(nogub), combos=len(combos),
                  nivmax=max(niveles) if niveles else 0)

mx_nac = max(r["dnac"] for r in raw.values())
mx_int = max(r["dint"] for r in raw.values())
mx_sis = max(r["sis"] for r in raw.values())
mx_com = max(r["combos"] for r in raw.values())
check(mx_nac == 3 and mx_int == 2 and mx_sis == 5 and mx_com == 2,
      f"máximos 2025 inesperados: nac={mx_nac} int={mx_int} sis={mx_sis} combos={mx_com}")

def v_cant(n):
    return 0 if n == 0 else 25 if n <= 3 else 50 if n <= 6 else 75 if n <= 9 else 100

py = {}
for p in PAISES:
    d = raw[p]
    if d["n"] == 0:
        py[p] = dict(s1L=0, s2=0, iL=0, v1=0, v2=0, v3=0, v4=0, v5=0, s1N=0, iN=0)
        continue
    vt = d["tipos"] / 5 * 100; ve = d["etapas"] / 4 * 100
    s1L_f = (vt + ve + d["cont"] / 2 * 100 + d["orgs"] / 3 * 100) / 4
    dev = (d["dnac"] / mx_nac * 100 + d["dint"] / mx_int * 100) / 2
    s2_f = (dev + d["sis"] / mx_sis * 100) / 2
    v3 = d["nivmax"] / 3 * 100
    v4 = (d["gub"] / 3 + d["nogub"] / 4 + d["combos"] / mx_com) / 3 * 100
    v5 = v_cant(d["n"])
    s1N_f = (vt + ve + v3 + v4 + v5) / 5
    py[p] = dict(s1L=rhu(s1L_f), s2=rhu(s2_f), iL=rhu((rhu(s1L_f) + rhu(s2_f)) / 2),
                 v1=vt, v2=ve, v3=v3, v4=v4, v5=v5, s1N=s1N_f,
                 iN=rhu((rhu(s1N_f) + rhu(s2_f)) / 2))

# --------------------------------------------------------------------------
# 2. Comparar contra las pestañas 2 y 3 (valores cacheados)
# --------------------------------------------------------------------------
w2 = wb["2 · Cálculo 2025 · f. antigua"]
w3 = wb["3 · Cálculo 2025 · f. nueva"]

def close(a, b, tol=0.05):
    if a is None: a = 0
    return abs(float(a) - float(b)) <= tol

for i, p in enumerate(PAISES):
    r = 3 + i
    e = py[p]
    check(close(w2.cell(r, 12).value, e["s1L"]), f"{p}: Sub1 legacy Excel={w2.cell(r,12).value} py={e['s1L']}")
    check(close(w2.cell(r, 19).value, e["s2"]), f"{p}: Sub2 Excel={w2.cell(r,19).value} py={e['s2']}")
    check(close(w2.cell(r, 20).value, e["iL"]), f"{p}: Ind legacy Excel={w2.cell(r,20).value} py={e['iL']}")
    for cidx, k in ((4, "v1"), (6, "v2"), (8, "v3"), (12, "v4"), (13, "v5"), (14, "s1N")):
        check(close(w3.cell(r, cidx).value, e[k]), f"{p}: {k} Excel={w3.cell(r,cidx).value} py={e[k]}")
    check(close(w3.cell(r, 20).value, e["iN"]), f"{p}: Ind nueva Excel={w3.cell(r,20).value} py={e['iN']}")

# --------------------------------------------------------------------------
# 3. Oficial 2025 (±1) y columna 2026 de la pestaña 4
# --------------------------------------------------------------------------
wb_b = openpyxl.load_workbook(BBDD, data_only=True)
oficial = {str(row[0]).strip(): row[3]
           for row in wb_b["Gráfico - Índice"].iter_rows(min_row=2, values_only=True) if row[0]}
n_ok = 0
for p in PAISES:
    if p not in oficial:
        continue
    check(abs(py[p]["iL"] - float(oficial[p])) <= 1,
          f"{p}: legacy {py[p]['iL']} vs oficial {oficial[p]} (>±1)")
    n_ok += 1
check(n_ok == 19, f"control oficial: {n_ok}/19 países")

wb1 = openpyxl.load_workbook(D1, data_only=True)
w26 = wb1["2 · Índice por país"]
ind26 = {w26.cell(r, 1).value: w26.cell(r, 23).value for r in range(3, 23)}
w4 = wb["4 · Comparación"]
for i, p in enumerate(PAISES):
    r = 6 + i
    check(close(w4.cell(r, 4).value, ind26[p]), f"{p}: 2026 pestaña4={w4.cell(r,4).value} vs D1={ind26[p]}")
    check(close(w4.cell(r, 5).value, py[p]["iN"] - py[p]["iL"]), f"{p}: Δ fórmula")
    check(close(w4.cell(r, 6).value, float(ind26[p]) - py[p]["iN"]), f"{p}: Δ real")

# --------------------------------------------------------------------------
# 4. Gráficos presentes
# --------------------------------------------------------------------------
with zipfile.ZipFile(OUT) as z:
    n_charts = sum(1 for n in z.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml"))
check(n_charts == 8, f"gráficos: {n_charts}/8")

# --------------------------------------------------------------------------
print(f"{'País':<5}{'IndLegacy':>10}{'Oficial':>9}{'IndNueva':>10}{'Ind2026':>9}")
for p in PAISES:
    print(f"{p:<5}{py[p]['iL']:>10}{str(oficial.get(p,'s/d')):>9}{py[p]['iN']:>10}{str(ind26[p]):>9}")
print(f"\nMáximos 2025: dev_nac={mx_nac} dev_int={mx_int} sistemas={mx_sis} combinaciones={mx_com} · gráficos={n_charts}")
if errors:
    print(f"\nRESULTADO: ✗ FALLÓ con {len(errors)} problema(s):")
    for e in errors:
        print("   -", e)
    sys.exit(1)
print("\nRESULTADO: ✓ TODO CUADRA (python ↔ Excel = 0 diferencias; oficial 2025 ±1 en 19/19; 2026 consistente con el entregable 1)")
