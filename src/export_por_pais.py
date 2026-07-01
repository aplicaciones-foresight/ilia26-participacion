#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export_por_pais.py — Planilla SIMPLE por país para validación con partners.

Un archivo con una pestaña por país. Cada fila = un caso, con lo mínimo para que
un partner valide: nombre, descripción breve, fuente (URL validada) y si se incluye.
Pensada para enviar por país; lo más simple posible.

Salida: data/gates/planilla_validacion_por_pais_ILIA2026.xlsx
"""
from __future__ import annotations
import glob, json, os, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import export_planilla_simple as eps  # reusa la clasificación (EXCLUIR, is_in, is_duda)

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

AZUL=PatternFill("solid",fgColor="1F4E78"); VERDE=PatternFill("solid",fgColor="C6EFCE")
ROJO=PatternFill("solid",fgColor="FFC7CE"); AMBAR=PatternFill("solid",fgColor="FFEB9C")
GRIS=PatternFill("solid",fgColor="D9D9D9")
WB=Font(bold=True,color="FFFFFF"); CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
TL=Alignment(horizontal="left",vertical="top",wrap_text=True)
TH=Side(style="thin",color="BFBFBF"); BOR=Border(left=TH,right=TH,top=TH,bottom=TH)

PAIS_NOMBRE={"AR":"Argentina","BO":"Bolivia","BR":"Brasil","CL":"Chile","CO":"Colombia",
 "CR":"Costa Rica","CU":"Cuba","DO":"Rep. Dominicana","EC":"Ecuador","GT":"Guatemala",
 "HN":"Honduras","JM":"Jamaica","MX":"México","PA":"Panamá","PE":"Perú","PY":"Paraguay",
 "SV":"El Salvador","TT":"Trinidad y Tobago","UY":"Uruguay","VE":"Venezuela","LATAM":"LATAM / Regional"}

URLRE=re.compile(r'^https?://', re.I)

# Universo ILIA: 20 países (aunque no tengan casos identificados este ciclo).
PAISES_ILIA=["AR","BO","BR","CL","CO","CR","CU","DO","EC","GT","HN","JM","MX","PA","PE","PY","SV","TT","UY","VE"]


def estado(d):
    cid=d["caso_id"]
    if cid in eps.EXCLUIR or not eps.is_in(d): return "NO"
    if eps.is_duda(d): return "DUDA"
    return "SÍ"


def motivo(d):
    cid=d["caso_id"]
    if cid in eps.EXCLUIR: return eps.EXCLUIR[cid]
    if not eps.is_in(d):
        return (d.get("motivo_exclusion_2025") or d.get("justificacion_elegibilidad") or "")[:200]
    if eps.is_duda(d): return "Pendiente de validación: confirmar si hay IA en el proceso."
    if d.get("cuenta_como_iniciativa") is False:
        return "Cuenta junto con su par (no se suma como iniciativa aparte)."
    return ""


def descripcion(d):
    p=(d.get("proceso_participativo") or d.get("justificacion_elegibilidad") or "").strip()
    p=re.sub(r'\s+',' ',p)
    if not p: return ""
    # primera(s) oración(es) hasta ~200 caracteres
    out=""
    for frag in re.split(r'(?<=[.…])\s', p):
        if not out: out=frag
        elif len(out)+len(frag)<=200: out+=" "+frag
        else: break
    if len(out)>230: out=out[:227]+"…"
    return out


def fuente(d):
    for s in (d.get("fuentes") or []):
        u=(s.get("url") or "").strip()
        if URLRE.match(u): return u
    return ""


def tab_name(pais):
    nom=PAIS_NOMBRE.get(pais,pais)
    n=f"{pais} - {nom}" if pais!="LATAM" else "LATAM - Regional"
    return n[:31]


def main():
    fichas=[json.load(open(f,encoding="utf-8")) for f in sorted(glob.glob(os.path.join(ROOT,"data/fichas/*_detalle.json")))]
    porpais={}
    for d in fichas: porpais.setdefault(str(d.get("pais")),[]).append(d)
    # los 20 países ILIA (con o sin casos) + LATAM/Regional al final
    paises=list(PAISES_ILIA) + (["LATAM"] if "LATAM" in porpais else [])
    sin_casos=[p for p in PAISES_ILIA if not porpais.get(p)]

    wb=Workbook(); wb.remove(wb.active)

    # ---------- Portada / índice ----------
    ws=wb.create_sheet("Instrucciones")
    filas=[
     ["VALIDACIÓN POR PAÍS — IA en Participación Ciudadana (ILIA 2026)"],
     ["",""],
     ["Qué es","Inventario de casos por país. Una pestaña por país. Por favor valide los casos de su país."],
     ["Qué revisar","1) ¿El caso EXISTE y está bien descrito?  2) ¿La fuente es correcta?  3) ¿Está bien la decisión de incluir (SÍ) o no (NO)?  4) ¿Falta algún caso?"],
     ["Cómo responder","Use la última columna «Comentario / corrección» para confirmar («OK») o corregir. Puede agregar filas al final con casos nuevos."],
     ["Criterio de inclusión","ENTRA si es un proceso de participación ciudadana (consulta, presupuesto participativo, mini-público, referendo, gobernanza colaborativa) que usa IA en alguna etapa (incluido el apoyo logístico). NO entra: civic tech de servicio/reporte, atención ciudadana, o voto/conteo electoral."],
     ["Significado de la columna «¿Se incluye?»","SÍ = entra al indicador · NO = queda fuera (ver Motivo) · DUDA = pendiente de validar."],
     ["",""],
     ["Cobertura","Los 20 países del ILIA tienen pestaña. Los que no tienen casos identificados este ciclo se marcan «SIN CASOS» (por favor confirme si conoce alguno)."],
     ["",""],
     ["ÍNDICE (20 países ILIA + Regional)","País — casos (de los cuales SÍ / DUDA / NO)"],
    ]
    for row in filas: ws.append(row)
    for p in paises:
        cs=porpais.get(p,[])
        if not cs:
            ws.append([f"{p} — {PAIS_NOMBRE.get(p,p)}", "SIN CASOS identificados este ciclo"]); continue
        nsi=sum(1 for d in cs if estado(d)=="SÍ"); nd=sum(1 for d in cs if estado(d)=="DUDA"); nno=sum(1 for d in cs if estado(d)=="NO")
        ws.append([f"{p} — {PAIS_NOMBRE.get(p,p)}", f"{len(cs)} casos  (SÍ {nsi} · DUDA {nd} · NO {nno})"])
    if sin_casos:
        ws.append(["",""])
        ws.append(["Países SIN casos", ", ".join(f"{p} ({PAIS_NOMBRE.get(p,p)})" for p in sin_casos)])
    ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=100
    for r in range(1,ws.max_row+1):
        ws.cell(row=r,column=1).font=Font(bold=True); ws.cell(row=r,column=1).alignment=TL; ws.cell(row=r,column=2).alignment=TL
    ws["A1"].font=Font(bold=True,size=13)

    # ---------- una pestaña por país ----------
    cols=["Caso","Descripción breve","Fuente (URL validada)","¿Se incluye?","Motivo (si NO/DUDA)","Comentario / corrección (partner)"]
    for p in paises:
        ws=wb.create_sheet(tab_name(p))
        ws.append(cols)
        casos_p=porpais.get(p,[])
        if not casos_p:
            ws.append(["(sin casos identificados este ciclo)",
                       "No se identificaron casos de IA en participación ciudadana para este país. Si conoce alguno, agréguelo en las filas de abajo.",
                       "", "—", "", ""])
        for d in sorted(casos_p, key=lambda x:(estado(x)!="SÍ", str(x.get("nombre_caso")))):
            ws.append([d.get("nombre_caso") or "", descripcion(d), fuente(d), estado(d), motivo(d), ""])
        # estilo cabecera
        for c in range(1,len(cols)+1):
            cell=ws.cell(row=1,column=c); cell.fill=AZUL; cell.font=WB; cell.alignment=CEN; cell.border=BOR
        ws.row_dimensions[1].height=30; ws.freeze_panes="A2"
        ws.auto_filter.ref=f"A1:{get_column_letter(len(cols))}{ws.max_row}"
        widths=[34,60,46,12,42,34]
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
        for r in range(2,ws.max_row+1):
            for c in range(1,len(cols)+1):
                ws.cell(row=r,column=c).alignment=TL; ws.cell(row=r,column=c).border=BOR
            inc=ws.cell(row=r,column=4); v=str(inc.value); inc.alignment=CEN; inc.font=Font(bold=True)
            inc.fill=VERDE if v=="SÍ" else ROJO if v=="NO" else AMBAR

    dst=os.path.join(ROOT,"data/gates/planilla_validacion_por_pais_ILIA2026.xlsx"); wb.save(dst)
    tot=len(fichas); nsi=sum(1 for d in fichas if estado(d)=="SÍ"); nd=sum(1 for d in fichas if estado(d)=="DUDA")
    print(f"[ok] {dst}")
    print(f"     {len(paises)} pestañas de país · {tot} casos (SÍ {nsi} · DUDA {nd} · NO {tot-nsi-nd})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
