#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_entregables.py — Verificación independiente de los entregables ILIA 2026
================================================================================
Recomputa TODO el índice CENIA v2 en PYTHON PURO directamente desde el coding de
la pestaña 1 del entregable 1 (SIN usar las fórmulas del Excel) y lo compara
contra los VALORES CACHEADOS de la pestaña 2 (tras recalc), con tolerancia 0.05.

Además verifica:
  · 28 casos incluidos · 80 excluidos con conteos 12/49/19
    (79 del insumo de validación + GeTAU AR, alta del equipo jul-2026).
  · Los 28 casos de la BBDD 2025 están todos contabilizados
    (16 siguen incluidos —e-Cidadania fusionado— y 12 en el grupo 1).
  · Las 49 exclusiones 2025 están todas en el grupo 2.

Sale con código ≠ 0 si algo no cuadra.
Uso:  python src/verify_entregables.py
"""
import sys, unicodedata
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

REPO = "/home/user/ilia26-participacion"
OUT1 = REPO + "/data/final/BBDD_casos_incluidos_ILIA2026.xlsx"
OUT2 = REPO + "/data/final/Casos_excluidos_ILIA2026.xlsx"
BASE = REPO + "/data/baseline/BBDD_IA_Participacion_2025.xlsx"

PAISES = ["AR","BO","BR","CL","CO","CR","CU","DO","EC","GT",
          "HN","JM","MX","PA","PE","PY","SV","TT","UY","VE"]
TP_TOKENS = ["participación digital","gobernanza colaborativa","mini públicos",
             "referendos e iniciativas populares","presupuestos participativos"]
CONV = ["gobierno local","gobierno nacional","otra institución pública","empresa",
        "universidad","sociedad civil","organización internacional"]
CODES = ["GL","GN","OIP","EMP","UNI","SC","OI"]
GUB = [0,1,2]; NOGUB = [3,4,5,6]
DEV = ["privado","academia","gobierno","sociedad civil"]
ANIO_REF, CAD = 2026, 2
TOL = 0.05
errors = []
def check(cond, msg):
    if not cond: errors.append(msg)
    return cond

def rhu(x):
    return int(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
def strip_ac(s):
    return "".join(c for c in unicodedata.normalize("NFKD", str(s)) if not unicodedata.combining(c))
def nk(s):
    s = strip_ac(str(s or "").lower())
    for ch in "\"'’“”().,:;/+-–—": s = s.replace(ch," ")
    return " ".join(s.split())
def norm(s):
    s = str(s or "").replace("; ",";").replace(", ",";").replace(",",";")
    return ";" + s.lower() + ";"
def cantidad(n):
    if n <= 0: return 0
    if n <= 3: return 25
    if n <= 6: return 50
    if n <= 9: return 75
    return 100

# --------------------------------------------------------------------------
# 1. Leer coding (literal) de la pestaña 1 del entregable 1
# --------------------------------------------------------------------------
wb = openpyxl.load_workbook(OUT1, data_only=True)
ws1 = wb["1 · BBDD Casos incluidos"]
# columnas por cabecera
hdr = {ws1.cell(1,c).value: c for c in range(1, ws1.max_column+1)}
def col(name):
    for k,v in hdr.items():
        if k and str(k).strip()==name: return v
    for k,v in hdr.items():
        if k and name.lower() in str(k).lower(): return v
    raise KeyError(name)
c_pais=col("País (ISO-2)"); c_cuenta=col("Cuenta como iniciativa"); c_tp=col("Tipo de proceso")
c_et=col("Etapas uso IA"); c_niv=col("Nivel (0-3)"); c_anio=col("Año"); c_conv=col("Convocante (7-cat)")
c_dev=col("Tipo desarrollador IA"); c_orig=col("Origen desarrollador"); c_tia=col("Tipos/familia IA")

casos=[]
for r in range(2, ws1.max_row+1):
    p = ws1.cell(r,c_pais).value
    if not p: continue
    if str(ws1.cell(r,c_caso if False else col("Caso")).value or "").strip()=="" : continue
    try: anio=int(float(ws1.cell(r,c_anio).value))
    except Exception: anio=ws1.cell(r,c_anio).value
    try: niv=int(float(ws1.cell(r,c_niv).value))
    except Exception: niv=0
    casos.append(dict(pais=str(p).strip(), cuenta=str(ws1.cell(r,c_cuenta).value or "").strip(),
        tp=ws1.cell(r,c_tp).value, et=ws1.cell(r,c_et).value, niv=niv, anio=anio,
        conv=ws1.cell(r,c_conv).value, dev=ws1.cell(r,c_dev).value,
        orig=ws1.cell(r,c_orig).value, tia=ws1.cell(r,c_tia).value))
check(len(casos)==28, f"Entregable 1: {len(casos)} casos incluidos (esperados 28)")

# catálogo de sistemas (mismo criterio que el build: excluye vacío y 'null')
catalog=[]
for c in casos:
    for t in str(c["tia"] or "").replace("; ",";").replace(", ",";").replace(",",";").split(";"):
        t=t.strip().lower()
        if t and t!="null" and t not in catalog: catalog.append(t)

# --------------------------------------------------------------------------
# 2. Recomputar flags y agregados por país (PYTHON PURO)
# --------------------------------------------------------------------------
def flags(c):
    ntp=norm(c["tp"]); net=norm(c["et"]); ncv=norm(c["conv"]); ndv=norm(c["dev"]); nor=norm(c["orig"])
    tp=[1 if f";{TP_TOKENS[i]};" in ntp else 0 for i in range(5)]
    et=[1 if f";{t};" in net else 0 for t in ("planificación","implementación","análisis")]
    et.append(1 if "traducción" in net else 0)
    cv=[1 if f";{CONV[i]};" in ncv else 0 for i in range(7)]
    dv=[1 if f";{DEV[i]};" in ndv else 0 for i in range(4)]
    onac=1 if ";nacional;" in nor else 0
    oint=1 if ";internacional;" in nor else 0
    nivef=0 if (c["niv"]==1 and (ANIO_REF-c["anio"])>CAD) else c["niv"]
    clave=""
    if sum(cv)>=2:
        clave="".join(CODES[i]+"+" for i in range(7) if cv[i]==1)
    sis=[1 if f";{tok};" in norm(c["tia"]) else 0 for tok in catalog]
    return dict(tp=tp,et=et,cv=cv,dv=dv,onac=onac,oint=oint,nivef=nivef,clave=clave,sis=sis)

for c in casos: c["fl"]=flags(c)

def agg(pais):
    cp=[c for c in casos if c["pais"]==pais]
    n=sum(1 for c in cp if c["cuenta"]=="sí")
    def anyflag(key,i): return 1 if any(c["fl"][key][i]==1 for c in cp) else 0
    ntp=sum(anyflag("tp",i) for i in range(5))
    net=sum(anyflag("et",i) for i in range(4))
    nivmax=max([c["fl"]["nivef"] for c in cp], default=0)
    ngub=sum(anyflag("cv",i) for i in GUB)
    nnog=sum(anyflag("cv",i) for i in NOGUB)
    combos=len({c["fl"]["clave"] for c in cp if c["fl"]["clave"]})
    devnac=sum(1 for i in range(4) if any(c["fl"]["dv"][i]==1 and c["fl"]["onac"]==1 for c in cp))
    devint=sum(1 for i in range(4) if any(c["fl"]["dv"][i]==1 and c["fl"]["oint"]==1 for c in cp))
    nsis=sum(1 for j in range(len(catalog)) if any(c["fl"]["sis"][j]==1 for c in cp))
    return dict(n=n,ntp=ntp,net=net,nivmax=nivmax,ngub=ngub,nnog=nnog,combos=combos,
                devnac=devnac,devint=devint,nsis=nsis)

A={p:agg(p) for p in PAISES}
maxcomb=max(A[p]["combos"] for p in PAISES)
maxnac=max(A[p]["devnac"] for p in PAISES)
maxint=max(A[p]["devint"] for p in PAISES)
maxsis=max(A[p]["nsis"] for p in PAISES)

def scores(p):
    a=A[p]
    v1=a["ntp"]/5*100; v2=a["net"]/4*100; v3=a["nivmax"]/3*100
    co=(a["combos"]/maxcomb) if maxcomb else 0
    v4=(a["ngub"]/3 + a["nnog"]/4 + co)/3*100
    v5=cantidad(a["n"])
    sub1=(v1+v2+v3+v4+v5)/5
    dn=(a["devnac"]/maxnac) if maxnac else 0
    di=(a["devint"]/maxint) if maxint else 0
    vdev=(dn+di)/2*100
    vsis=(a["nsis"]/maxsis*100) if maxsis else 0
    sub2=(vdev+vsis)/2
    s1r=rhu(sub1); s2r=rhu(sub2); indf=rhu((s1r+s2r)/2); indr=(sub1+sub2)/2
    return dict(n=a["n"],ntp=a["ntp"],v1=v1,net=a["net"],v2=v2,nivmax=a["nivmax"],v3=v3,
                ngub=a["ngub"],nnog=a["nnog"],combos=a["combos"],v4=v4,v5=v5,sub1=sub1,
                devnac=a["devnac"],devint=a["devint"],vdev=vdev,nsis=a["nsis"],vtia=vsis,
                sub2=sub2,s1r=s1r,s2r=s2r,indf=indf,indr=indr)
S={p:scores(p) for p in PAISES}

# --------------------------------------------------------------------------
# 3. Leer valores CACHEADOS de la pestaña 2 y comparar
# --------------------------------------------------------------------------
ix=wb["2 · Índice por país"]
M=dict(n="B",ntp="C",v1="D",net="E",v2="F",nivmax="G",v3="H",ngub="I",nnog="J",
       combos="K",v4="L",v5="M",sub1="N",devnac="O",devint="P",vdev="Q",nsis="R",
       vtia="S",sub2="T",s1r="U",s2r="V",indf="W",indr="X")
MR0=3
cmp_keys=["n","ntp","v1","net","v2","nivmax","v3","ngub","nnog","combos","v4","v5",
          "sub1","devnac","devint","vdev","nsis","vtia","sub2","s1r","s2r","indf","indr"]
mismatches=0
report=[]
for i,p in enumerate(PAISES):
    r=MR0+i
    excel={}
    for k in cmp_keys:
        v=ix[f"{M[k]}{r}"].value
        excel[k]=float(v) if v is not None else None
    for k in cmp_keys:
        py=float(S[p][k]); ex=excel[k]
        if ex is None or abs(py-ex)>TOL:
            mismatches+=1
            errors.append(f"[{p}] {k}: python={py:.4f} excel={ex}")
    report.append((p,S[p]))
check(mismatches==0, f"{mismatches} discrepancias python vs Excel (pestaña 2)")

# --------------------------------------------------------------------------
# 4. Máximos relativos esperados (control)
# --------------------------------------------------------------------------
check(maxcomb==2, f"máx combinaciones = {maxcomb} (esperado 2)")
check(maxnac==4, f"máx dev nacional = {maxnac} (esperado 4)")
check(maxint==3, f"máx dev internacional = {maxint} (esperado 3)")
check(maxsis==11, f"máx sistemas = {maxsis} (esperado 11)")
check(len(catalog)==25, f"catálogo de sistemas = {len(catalog)} tokens (esperado 25)")
# N de control
for p,exp in {"BR":7,"CL":9,"CO":3}.items():
    check(A[p]["n"]==exp, f"N iniciativas {p} = {A[p]['n']} (esperado {exp})")
for p in PAISES:
    if p not in ("BR","CL","CO"):
        casos_p=[c for c in casos if c["pais"]==p]
        if casos_p:
            check(A[p]["n"]==1, f"N iniciativas {p} = {A[p]['n']} (esperado 1)")

# --------------------------------------------------------------------------
# 5. Entregable 2: conteos y conciliación con BBDD 2025
# --------------------------------------------------------------------------
wb2=openpyxl.load_workbook(OUT2, data_only=True)
ws2=wb2["Casos excluidos"]
# encontrar fila de cabecera
hrow=None
for r in range(1,12):
    if str(ws2.cell(r,1).value or "").strip()=="País": hrow=r; break
check(hrow is not None, "No se encontró la cabecera de la pestaña de excluidos")
exrows=[]
for r in range(hrow+1, ws2.max_row+1):
    p=ws2.cell(r,1).value
    if not p: continue
    exrows.append(dict(pais=str(p).strip(), caso=str(ws2.cell(r,2).value or "").strip(),
                       origen=str(ws2.cell(r,3).value or "").strip(),
                       det=str(ws2.cell(r,4).value or ""), motivo=str(ws2.cell(r,6).value or "")))
check(len(exrows)==80, f"Entregable 2: {len(exrows)} filas (esperadas 80: 79 del insumo + GeTAU)")

# Guardas anti-truncamiento y de saneo (hallazgo del QA externo 2026-07-20:
# el insumo traía 14 motivos capados a 300 chars — reconstruidos con etiqueta).
import re as _re
_FIN_OK = ".»\")!?]"
for e in exrows:
    m, d = e["motivo"], e["det"]
    check(len(m) >= 10, f"motivo vacío: {e['pais']} {e['caso'][:40]}")
    check(not (295 <= len(m) <= 305 and (not m or m[-1] not in _FIN_OK)),
          f"motivo con firma de truncado a ~300: {e['pais']} {e['caso'][:40]} (fin={m[-25:]!r})")
    check(not _re.search(r"https?://\S{0,14}$", m),
          f"motivo termina en URL rota: {e['pais']} {e['caso'][:40]}")
    check(not any(ch in m + d for ch in "​‌‍﻿"),
          f"caracteres de ancho cero: {e['pais']} {e['caso'][:40]}")
    if e["origen"].startswith("Nuevo 2026"):
        check("Excluido en 2025 por" not in d,
              f"detalle-plantilla 2025 en caso Nuevo 2026: {e['pais']} {e['caso'][:40]}")
_n_rec = sum(1 for e in exrows if "[Cierre reconstruido" in e["motivo"])
check(_n_rec == 14, f"motivos con etiqueta de reconstrucción = {_n_rec} (esperados 14)")
def grp(origen):
    if origen.startswith("Base 2025 — estaba INCLUIDO"): return 1
    if origen.startswith("Base 2025 — ya estaba EXCLUIDO"): return 2
    if origen.startswith("Nuevo 2026"): return 3
    return 0
gc={1:0,2:0,3:0,0:0}
for e in exrows: gc[grp(e["origen"])]+=1
check(gc[1]==12, f"grupo 1 = {gc[1]} (esperado 12)")
check(gc[2]==49, f"grupo 2 = {gc[2]} (esperado 49)")
check(gc[3]==19, f"grupo 3 = {gc[3]} (esperado 19: 18 del insumo + GeTAU)")
check(gc[0]==0, f"{gc[0]} filas con origen no reconocido")

g1=[e for e in exrows if grp(e["origen"])==1]
g2=[e for e in exrows if grp(e["origen"])==2]

# BBDD 2025
wbb=openpyxl.load_workbook(BASE, data_only=True)
bc=wbb["BBDD Casos"]; be=wbb["BBDD Casos Excluidos"]
bbdd_casos=[]
for r in range(2, bc.max_row+1):
    p=bc.cell(r,1).value; caso=bc.cell(r,2).value
    if not p or not caso: continue
    bbdd_casos.append(dict(pais=str(p).strip(), caso=str(caso).strip(), _row=r))
bbdd_excl=[]
for r in range(2, be.max_row+1):
    p=be.cell(r,2).value; caso=be.cell(r,3).value
    if not p or not caso: continue
    bbdd_excl.append(dict(pais=str(p).strip(), caso=str(caso).strip()))
check(len(bbdd_casos)==28, f"BBDD 2025 casos = {len(bbdd_casos)} (esperado 28)")
check(len(bbdd_excl)==49, f"BBDD 2025 excluidos = {len(bbdd_excl)} (esperado 49)")

incl_set=[dict(pais=c["pais"], caso=None) for c in casos]  # placeholder
# nombres del entregable 1 (incluidos)
incl_names=[(ws1.cell(r,c_pais).value, ws1.cell(r,col("Caso")).value)
            for r in range(2, ws1.max_row+1) if ws1.cell(r,c_pais).value]
incl_names=[(str(p).strip(), nk(caso)) for p,caso in incl_names]
g1_names=[(e["pais"], nk(e["caso"])) for e in g1]
g2_names=[(e["pais"], nk(e["caso"])) for e in g2]

# Conciliación DETERMINÍSTICA: igualdad exacta de nombre normalizado (nk),
# más un mapa explícito para los nombres que cambiaron entre 2025 y 2026.
# (El matching difuso por solapamiento de palabras producía falsos positivos:
#  «ParticipACT Brasil» ≠ «Brasil Participativo - clustering semántico…».)

# 2025 → nombre en el entregable (cuando NO es igualdad exacta)
ALIAS_2025 = {
    # fusión: el caso 2025 de e-Cidadania quedó FUSIONADO en la fila 2026
    ("BR", nk("Portal e-Cidadania: Marcado automático de respuestas ciudadanas en videos de audiencias públicas")):
        ("INCLUIDO", nk("Portal e-Cidadania (Senado Federal de Brasil) — IA sobre aportes ciudadanos: marcado automático de respuestas en audiencias públicas + matching semántico de ideias legislativas")),
    # exclusiones 2025 cuyo título cambió levemente en la planilla validada
    ("CL", nk("Plataforma UCampus – participación 2023")):
        ("G2", nk("Plataforma UCampus - participación 2023 (proceso constitucional)")),
    ("CO", nk("\"Pa’ que veás”")):
        ("G2", nk("Pa' que veás (Cali)")),
    ("CO", nk("iLabs – Laboratorios de Inteligencia Colectiva")):
        ("G2", nk("iLabs - Laboratorios de Inteligencia Colectiva (Ideemos)")),
    ("BR", nk("De ciudadano a senador: la inteligencia artificial y la reinvención de la elaboración legislativa en Brasil")):
        ("G2", nk("De ciudadano a senador: la inteligencia artificial y la reinvención de la elaboración legislativa en Brasil”")),
}

incl_set = set(incl_names)
g1_set = set(g1_names)
g2_set = set(g2_names)

def resolver(pais, caso_2025):
    """Devuelve 'INCLUIDO', 'G1' o 'G2' para un caso de la BBDD 2025, o None."""
    key = (pais, nk(caso_2025))
    if key in ALIAS_2025:
        destino, alias = ALIAS_2025[key]
        objetivo = {"INCLUIDO": incl_set, "G1": g1_set, "G2": g2_set}[destino]
        return destino if (pais, alias) in objetivo else None
    en_i = key in incl_set
    en_1 = key in g1_set
    en_2 = key in g2_set
    if en_i and en_1:  # homónimo (CL Proceso Constitucional): 1ª aparición → incluido, 2ª → G1
        return "HOMO"
    if en_i: return "INCLUIDO"
    if en_1: return "G1"
    if en_2: return "G2"
    return None

# conciliar 28 BBDD casos → incluido 2026 o grupo 1
homo_used={}
n_incl=0; n_g1=0; unmatched=[]
for c in sorted(bbdd_casos, key=lambda x:x["_row"]):
    r = resolver(c["pais"], c["caso"])
    if r == "HOMO":
        key=(c["pais"], nk(c["caso"]))
        homo_used[key]=homo_used.get(key,0)+1
        if homo_used[key]==1: n_incl+=1
        else: n_g1+=1
    elif r == "INCLUIDO": n_incl+=1
    elif r == "G1": n_g1+=1
    else: unmatched.append((c["pais"],c["caso"]))
check(n_incl==16, f"BBDD 2025 casos que siguen incluidos = {n_incl} (esperado 16)")
check(n_g1==12, f"BBDD 2025 casos que pasan al grupo 1 = {n_g1} (esperado 12)")
check(not unmatched, f"BBDD 2025 casos sin conciliar: {unmatched}")

# conciliar 49 BBDD excluidos → grupo 2
n_g2=0; unmatched2=[]
for e in bbdd_excl:
    if resolver(e["pais"], e["caso"]) == "G2": n_g2+=1
    else: unmatched2.append((e["pais"],e["caso"]))
check(n_g2==49, f"BBDD 2025 exclusiones en grupo 2 = {n_g2} (esperado 49)")
check(not unmatched2, f"BBDD 2025 exclusiones sin conciliar: {unmatched2}")

# --------------------------------------------------------------------------
# 6. Reporte
# --------------------------------------------------------------------------
print("="*92)
print("VERIFICACIÓN ENTREGABLES ILIA 2026 — recómputo independiente (python) vs Excel (cacheado)")
print("="*92)
print(f"Casos incluidos: {len(casos)} · catálogo sistemas: {len(catalog)} tokens")
print(f"Máximos relativos del ciclo: combinaciones={maxcomb}  dev_nac={maxnac}  dev_int={maxint}  sistemas={maxsis}")
print("-"*92)
print(f"{'País':<5}{'N':>3}{'V1':>7}{'V2':>7}{'V3':>7}{'V4':>7}{'V5':>6}{'Sub1':>7}"
      f"{'Vdev':>7}{'Vsis':>7}{'Sub2':>7}{'S1r':>5}{'S2r':>5}{'Ind':>5}{'Ind1d':>7}")
for p in PAISES:
    s=S[p]
    print(f"{p:<5}{s['n']:>3}{s['v1']:>7.1f}{s['v2']:>7.1f}{s['v3']:>7.1f}{s['v4']:>7.1f}"
          f"{s['v5']:>6.0f}{s['sub1']:>7.1f}{s['vdev']:>7.1f}{s['vtia']:>7.1f}{s['sub2']:>7.1f}"
          f"{s['s1r']:>5}{s['s2r']:>5}{s['indf']:>5}{s['indr']:>7.1f}")
print("-"*92)
print(f"Entregable 2 — conteos por origen: G1={gc[1]}  G2={gc[2]}  G3={gc[3]}  total={sum(gc[k] for k in (1,2,3))}")
print(f"Conciliación BBDD 2025: incluidos-que-siguen={n_incl}/16  ·  pasan-a-grupo1={n_g1}/12  ·  exclusiones-en-grupo2={n_g2}/49")
print("-"*92)
if errors:
    print(f"RESULTADO: ✗ FALLÓ con {len(errors)} problema(s):")
    for e in errors[:60]: print("   -", e)
    sys.exit(1)
print("RESULTADO: ✓ TODO CUADRA (discrepancias python↔Excel = 0; conciliaciones OK)")
sys.exit(0)
