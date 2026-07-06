#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
calc_comparativo.py — Índice preliminar 2026 + comparación 2025↔2026 con la
MISMA metodología (CENIA v2). Aísla el efecto del cambio de datos (no de fórmula).

  · Set 2026 = casos con clasificación final SÍ (34 iniciativas efectivas).
  · Set 2025 = casos del baseline 2025 (bloque '2-baseline', como estaban en 2025).
  Ambos se corren con compute_2026 (metodología nueva).

Salida: data/gates/calculo_indice_2025vs2026_ILIA2026.xlsx
"""
from __future__ import annotations
import glob, json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import calc_engine as ce
import calc_preliminar as cp            # ficha_a_caso, cargar_config, PAIS_NOMBRE
import export_planilla_simple as eps    # clasificación final

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
AZUL=PatternFill("solid",fgColor="1F4E78"); VERDE=PatternFill("solid",fgColor="C6EFCE")
ROJO=PatternFill("solid",fgColor="FFC7CE"); AMBARH=PatternFill("solid",fgColor="FFF2CC")
WB=Font(bold=True,color="FFFFFF"); CEN=Alignment(horizontal="center",vertical="center",wrap_text=True)
TL=Alignment(horizontal="left",vertical="top",wrap_text=True)
TH=Side(style="thin",color="BFBFBF"); BOR=Border(left=TH,right=TH,top=TH,bottom=TH)
PN=cp.PAIS_NOMBRE


def estado(d):
    cid=d["caso_id"]
    if cid in eps.EXCLUIR or not eps.is_in(d): return "NO"
    if cid in eps.RESCATE_ENTRA: return "SÍ"
    if cid in eps.RECLASIF_DUDA or cid in eps.DUDA_REVISAR: return "DUDA"
    if eps.is_duda(d): return "DUDA"
    return "SÍ"


def _corre(fichas_set, cfg):
    casos=[cp.ficha_a_caso(d) for d in fichas_set]
    s1,s2,ind,_=ce.compute_2026(casos,cfg)
    return s1,s2,ind


def _prom(ind):
    reg=round(sum(ind[p] for p in ce.PAISES)/len(ce.PAISES))
    con=[p for p in ce.PAISES if ind[p]>0]
    return reg,(round(sum(ind[p] for p in con)/len(con)) if con else 0)


def hdr(ws,n,h=28):
    for c in range(1,n+1):
        cell=ws.cell(row=1,column=c); cell.fill=AZUL; cell.font=WB; cell.alignment=CEN; cell.border=BOR
    ws.row_dimensions[1].height=h; ws.freeze_panes="A2"

def widths(ws,wl):
    for i,w in enumerate(wl,1): ws.column_dimensions[get_column_letter(i)].width=w

def body(ws,ncol):
    for r in range(2,ws.max_row+1):
        for c in range(1,ncol+1):
            ws.cell(row=r,column=c).alignment=TL; ws.cell(row=r,column=c).border=BOR


def main():
    cfg=cp.cargar_config()
    fichas=[json.load(open(f,encoding="utf-8")) for f in sorted(glob.glob(os.path.join(ROOT,"data/fichas/*_detalle.json")))]
    def cuenta(d): return d.get("cuenta_como_iniciativa") is not False

    set26=[d for d in fichas if estado(d)=="SÍ" and cuenta(d)]
    # Set 2025: fichas del baseline TAL COMO ERAN en 2025. Si una ficha fue fusionada
    # después (p.ej. e-Cidadania consolidó atributos de una herramienta 2025/26), usa
    # los atributos originales guardados en 'baseline_2025_atributos'.
    set25=[]
    for d in fichas:
        if d.get("bloque")=="2-baseline" and cuenta(d):
            ov=d.get("baseline_2025_atributos")
            set25.append({**d,**ov} if ov else d)
    n26,n25=len(set26),len(set25)

    s1_26,s2_26,ind26=_corre(set26,cfg)
    _,_,ind25=_corre(set25,cfg)

    # nº iniciativas por país
    from collections import Counter
    c26=Counter(d.get("pais") for d in set26); c25=Counter(d.get("pais") for d in set25)
    con26=[p for p in ce.PAISES if s1_26[p]["n"]>0]

    wb=Workbook(); wb.remove(wb.active)

    # ---------- 1 · Índice 2026 (preliminar) ----------
    ws=wb.create_sheet("1 · Índice 2026")
    ws.append(["País","","Sub-Uso","Sub-Desarrollo","INDICADOR 2026","Nº iniciativas"])
    for p in sorted(con26,key=lambda q:-ind26[q]):
        ws.append([p,PN[p],round(s1_26[p]["v_pub"]),round(s2_26[p]["v_pub"]),ind26[p],c26[p]])
    pr,pc=_prom(ind26)
    ws.append([]); ws.append(["","Promedio (países con iniciativas)","","",pc,n26])
    ws.append(["","Promedio regional (20 países)","","",pr,""])
    hdr(ws,6); widths(ws,[6,40,12,15,15,14]); body(ws,6)
    for r in range(2,ws.max_row+1):
        for c in (3,4,5,6): ws.cell(row=r,column=c).alignment=CEN
        iv=ws.cell(row=r,column=5)
        if isinstance(iv.value,int):
            iv.font=Font(bold=True)
            if iv.value>=50: iv.fill=VERDE
            elif iv.value>=25: iv.fill=AMBARH

    # ---------- 2 · Comparación 2025 vs 2026 (misma metodología) ----------
    ws2=wb.create_sheet("2 · 2025 vs 2026 (misma met.)")
    ws2.append(["País","","Indicador 2025","Indicador 2026","Δ (2026−2025)","Nº inic. 2025","Nº inic. 2026"])
    orden=[p for p in ce.PAISES if ind25[p]>0 or ind26[p]>0]
    orden.sort(key=lambda q:-max(ind25[q],ind26[q]))
    for p in orden:
        ws2.append([p,PN[p],ind25[p],ind26[p],ind26[p]-ind25[p],c25.get(p,0),c26.get(p,0)])
    pr25,pc25=_prom(ind25); pr26,pc26=_prom(ind26)
    ws2.append([])
    ws2.append(["","Promedio regional (20 países)",pr25,pr26,pr26-pr25,n25,n26])
    ws2.append(["","Promedio (países con iniciativas)",pc25,pc26,pc26-pc25,"",""])
    hdr(ws2,7); widths(ws2,[6,36,15,15,16,14,14]); body(ws2,7)
    for r in range(2,ws2.max_row+1):
        for c in range(3,8): ws2.cell(row=r,column=c).alignment=CEN
        dv=ws2.cell(row=r,column=5)
        if isinstance(dv.value,int) and dv.value!=0:
            dv.font=Font(bold=True); dv.fill=VERDE if dv.value>0 else ROJO
            if dv.value>0: dv.value=f"+{dv.value}"

    # ---------- 3 · Sub-Uso 2026 (5 variables) ----------
    ws3=wb.create_sheet("3 · Sub-Uso 2026 (detalle)")
    ws3.append(["País","","Tipos proc.","Etapas","Continuidad","Org. Convocante","Cantidad","Sub-Uso","Nº inic.","Combos co-conv."])
    for p in sorted(con26,key=lambda q:-s1_26[q]["v"]):
        d=s1_26[p]
        ws3.append([p,PN[p],round(d["tipos"]),round(d["etapas"]),round(d["cont"]),round(d["conv"]),
                    round(d["cant"]),round(d["v"]),d["n"],d.get("co_combos","")])
    hdr(ws3,10); widths(ws3,[6,30,11,9,12,15,10,10,9,14]); body(ws3,10)
    for r in range(2,ws3.max_row+1):
        for c in range(3,11): ws3.cell(row=r,column=c).alignment=CEN

    # ---------- 4 · Supuestos ----------
    ws4=wb.create_sheet("4 · Supuestos y método")
    filas=[
     ["CÁLCULO DEL ÍNDICE — 2026 preliminar y comparación 2025 vs 2026"],["",""],
     ["Qué es","Cálculo preliminar y determinístico con los datos actuales. La comparación usa la MISMA metodología (CENIA v2) sobre ambas bases, para aislar el efecto del cambio de datos (no de fórmula)."],
     ["Set 2026",f"{n26} iniciativas efectivas: clasificación final SÍ (tras validación manual), respetando el dedup."],
     ["Set 2025",f"{n25} iniciativas del baseline 2025 (como estaban en el índice 2025), recodificadas a la taxonomía 2026 y corridas con la metodología nueva."],
     ["Metodología","CENIA v2 (jun-2026). Indicador = (Sub-Uso + Sub-Desarrollo)/2. Sub-Uso = promedio de 5: Tipos de proceso (÷5) · Etapas (÷4) · Continuidad = nivel máx. del país (÷3) · Organización Convocante (3 sub-componentes: gub ÷3 + no-gub ÷4 + co-convocatoria por máx. relativo) · Cantidad = todas las elegibles (umbrales 0/25/50/75/100). Sub-Desarrollo = 2025 (desarrollador nac/int × 4 tipos, máx. relativo + tipos de IA)."],
     ["Redondeo",f"{cfg.get('redondeo',{}).get('modo','legacy')} · count_min_level={cfg.get('cantidad',{}).get('count_min_level',0)} · escenario={cfg.get('escenarios',{}).get('scenario_activo','A')}"],
     ["Nota máx. relativo","El Sub-Desarrollo y la co-convocatoria usan máximo relativo POR CICLO: cada base (2025 y 2026) se normaliza contra su propio máximo, como indica la metodología."],
     ["Aviso","Preliminar. Es una foto con los datos de hoy; los valores pueden cambiar si se completan gaps o se ajustan casos frontera."],
    ]
    for row in filas: ws4.append(row)
    ws4.column_dimensions["A"].width=24; ws4.column_dimensions["B"].width=118
    for r in range(1,ws4.max_row+1):
        ws4.cell(row=r,column=1).font=Font(bold=True); ws4.cell(row=r,column=1).alignment=TL; ws4.cell(row=r,column=2).alignment=TL
    ws4["A1"].font=Font(bold=True,size=12)

    dst=os.path.join(ROOT,"data/gates/calculo_indice_2025vs2026_ILIA2026.xlsx"); wb.save(dst)
    print(f"[ok] {dst}")
    print(f"     2026: {n26} inic · prom regional={pr26} · prom con casos={pc26}")
    print(f"     2025 (misma met.): {n25} inic · prom regional={pr25} · prom con casos={pc25}")
    print(f"     Δ regional (misma metodología): {pr26-pr25:+d}")
    print("\n     País  Ind2025  Ind2026   Δ")
    for p in orden:
        print(f"     {p:>3}   {ind25[p]:>6}   {ind26[p]:>6}   {ind26[p]-ind25[p]:+d}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
