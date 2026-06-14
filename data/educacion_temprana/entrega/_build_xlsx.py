#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera el informe final en Excel (.xlsx) a partir de los CSV ya verificados.
Hojas: 'Resumen' (por país, ambos escenarios), 'Evidencia verbatim' (1 fila por cita,
con verbatim + URL + página) y 'Metodología' (rúbrica, regla adoptada, nota de electivos,
definición de secundaria con paralelo a Chile, verificación).
Uso: python data/educacion_temprana/entrega/_build_xlsx.py"""
import csv, pathlib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = pathlib.Path(__file__).resolve().parents[1]
OUT = BASE / "entrega" / "INFORME_FINAL_2026-06-06.xlsx"
def read_csv(name):
    with open(BASE / name, encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))
resumen = read_csv("calculo_comparado.csv")
evid = read_csv("informe_comparado.csv")

HDR_FILL = PatternFill("solid", fgColor="0B3D6B")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE = Font(bold=True, color="0B3D6B", size=14)
SUB = Font(bold=True, color="0B3D6B", size=11)
G5 = PatternFill("solid", fgColor="C6EFCE")   # verde cat 5
A4 = PatternFill("solid", fgColor="FFEB9C")   # ámbar cat 4
ALT = PatternFill("solid", fgColor="F4F7FB")  # fila par
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
thin = Side(style="thin", color="DDDDDD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def write_table(ws, rows, spec, cat_cols=(), url_col=None):
    ws.append([h for _, h, _, _ in spec])
    for r in rows:
        ws.append([r.get(k, "") for k, _, _, _ in spec])
    n = len(spec)
    for c in range(1, n + 1):
        cell = ws.cell(1, c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(n)}1"
    for i, (_, _, w, _) in enumerate(spec, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for ri, r in enumerate(rows, 2):
        for i, (k, _, _, wrap) in enumerate(spec, 1):
            cell = ws.cell(ri, i)
            cell.alignment = WRAP if wrap else TOP
            cell.font = Font(size=9)
            cell.border = BORDER
            if ri % 2 == 0:
                cell.fill = ALT
        for k in cat_cols:
            idx = next((j for j, (kk, _, _, _) in enumerate(spec, 1) if kk == k), None)
            if idx:
                v = str(r.get(k, "")).strip()
                if v == "5": ws.cell(ri, idx).fill = G5
                elif v == "4": ws.cell(ri, idx).fill = A4
        if url_col:
            idx = next((j for j, (kk, _, _, _) in enumerate(spec, 1) if kk == url_col), None)
            u = str(r.get(url_col, "")).strip()
            if idx and u.startswith("http"):
                cell = ws.cell(ri, idx)
                cell.hyperlink = u
                cell.font = Font(size=9, color="0563C1", underline="single")

wb = openpyxl.Workbook()

# ---------- Hoja 1: Resumen ----------
ws1 = wb.active; ws1.title = "Resumen"
spec1 = [("pais","País",16,False),("iso2","ISO",6,False),
         ("nivel_secundaria","Nivel de secundaria (local)",34,True),
         ("equivalente_chile","Equivalente en Chile",20,True),
         ("cat_inclusiva","Cat. inclusiva (ILIA)",11,False),("puntaje_inclusiva","Puntaje incl.",10,False),
         ("cat_estricta","Cat. estricta (informativa)",12,False),("puntaje_estricta","Puntaje estr.",10,False),
         ("argumento_inclusiva","Argumento (inclusiva)",52,True),
         ("argumento_estricta","Argumento (estricta, informativa)",42,True),
         ("nota_curso_electivo","Nota — curso electivo",46,True)]
write_table(ws1, resumen, spec1, cat_cols=("cat_inclusiva","cat_estricta"))

# ---------- Hoja 2: Evidencia verbatim ----------
ws2 = wb.create_sheet("Evidencia verbatim")
spec2 = [("pais","País",14,False),("iso2","ISO",6,False),
         ("cat_inclusiva","Cat. incl.",8,False),("puntaje_inclusiva","Pje incl.",8,False),
         ("cat_estricta","Cat. estr.",8,False),("puntaje_estricta","Pje estr.",8,False),
         ("argumento_inclusiva","Argumento (inclusiva)",44,True),
         ("argumento_estricta","Argumento (estricta, informativa)",38,True),
         ("nota_curso_electivo","Nota — curso electivo",40,True),
         ("tipo","Tipo",13,False),("nivel_secundaria","Nivel secundaria",26,True),
         ("equivalente_chile","Equiv. Chile",16,True),("documento","Documento oficial",38,True),
         ("ubicacion","Ubicación en el documento",40,True),("pagina_pdf","Pág. PDF",13,True),
         ("cita_verbatim","Cita verbatim",58,True),("traduccion_es","Traducción (ES)",46,True),
         ("url_fuente","URL fuente",40,True),("archivo_local","Archivo local",34,True),
         ("verificacion","Verificación",16,False)]
write_table(ws2, evid, spec2, cat_cols=("cat_inclusiva","cat_estricta"), url_col="url_fuente")

# ---------- Hoja 3: Metodología ----------
ws3 = wb.create_sheet("Metodología")
ws3.column_dimensions["A"].width = 26; ws3.column_dimensions["B"].width = 90
def put(row, a, b, a_font=None, b_font=None, wrap=True):
    ws3.cell(row, 1, a); ws3.cell(row, 2, b)
    if a_font: ws3.cell(row, 1).font = a_font
    if b_font: ws3.cell(row, 2).font = b_font
    ws3.cell(row, 1).alignment = Alignment(vertical="top", wrap_text=True)
    ws3.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=wrap)
ws3.cell(1, 1, "Educación temprana en IA — Subindicador (b) del ILIA · Resultados finales").font = TITLE
r = 3
put(r, "Encargo", "Replicación del subindicador (b) «Educación temprana en IA» del ILIA 2025 (CEPAL/CENIA) para 5 países benchmark no latinoamericanos (PT, ES, EE, DE, SG), por revisión del currículo escolar oficial vigente. El puntaje se asigna sobre secundaria."); r += 1
put(r, "Resultado", "ES = DE = EE = SG = 5 (100) · PT = 4 (75). Promedio del grupo: 95/100 (regla inclusiva)."); r += 1
put(r, "Regla adoptada", "INCLUSIVA, confirmada con el equipo del ILIA: los cursos electivos se consideran parte del currículo nacional; la IA en cualquier asignatura vigente (incl. electivas) cuenta como «IA implementada» (cat. 5)."); r += 1
put(r, "Nota — electivos (EE/SG)", "En Estonia y Singapur la IA del currículo nacional vigente se imparte en una asignatura ELECTIVA (Estonia: Informaatika; Singapur: Computing 7155) y como subtema; por la regla del ILIA, ambos = categoría 5. Se documenta por transparencia."); r += 1
put(r, "Sensibilidad (informativa)", "No adoptada. Una lectura estricta (contenido troncal/sustantivo) dejaría a EE y SG en 4 → promedio del grupo 85."); r += 1
put(r, "Verificación", "15/15 citas verificadas como subcadena exacta del documento oficial primario (fuentes/_verify_report.json). Fecha de verbatim: 2026-06-06."); r += 2

ws3.cell(r, 1, "Rúbrica (Cuadro 2, p. 63 del ILIA)").font = SUB; r += 1
for cat, desc, pje in [("1","No tiene propuesta","0"),("2","Propuesta TIC","25"),("3","Propuesta IA","50"),
                       ("4","TIC implementado","75"),("5","IA implementado","100")]:
    ws3.cell(r,1,f"Categoría {cat}"); ws3.cell(r,2,f"{desc}  →  {pje} puntos"); r += 1
r += 1
ws3.cell(r, 1, "Definición de «secundaria» y paralelo a Chile").font = SUB; r += 1
for pais, nivel, chile in [
    ("Portugal","3.º ciclo (7.º–9.º) + Ensino secundário (10.º–12.º)","7º básico – 4º medio"),
    ("España","ESO (1.º–4.º) + Bachillerato (1.º–2.º)","7º básico – 4º medio"),
    ("Estonia","Põhikool 7.º–9.º + Gümnaasium (10.º–12.º)","8º básico – 4º medio"),
    ("Alemania","Sek I (Kl. 5–10) + Sek II (Kl. 11–12/13)","5º básico – 4º medio"),
    ("Singapur","Secondary 1–4/5","8º básico – 3º medio")]:
    ws3.cell(r,1,pais); ws3.cell(r,2,f"{nivel}   ·   ≈ Chile: {chile}"); r += 1

wb.save(OUT)
print(f"OK -> {OUT.relative_to(BASE.parents[1])}  ({OUT.stat().st_size} bytes)")
print(f"Hojas: {wb.sheetnames}  | Resumen={len(resumen)} filas, Evidencia={len(evid)} filas")
