#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Entregables POR PAÍS, desde la fuente única (ECON/ECON_CELLS + matriz G):
  1) Una planilla .xlsx por país en por_pais/  -> para cada indicador: lo que hay + lo que falta/cómo obtenerlo.
  2) Resumen_por_pais_ILIA2026.pdf            -> por país: qué se hizo y qué queda.
Reusa datos_gobernanza.py y generar_pdf_indicadores.py (no recolecta nada nuevo)."""
import os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer, HRFlowable, KeepTogether, PageBreak)
from generar_pdf_indicadores import ECON, ECON_CELLS
from datos_gobernanza import G, st, GOB_META, COUNTRY_GIDX

HERE = os.path.dirname(os.path.abspath(__file__))
os.makedirs(os.path.join(HERE, "por_pais"), exist_ok=True)
ORDER = ["ES", "EE", "SG", "DE", "PT"]
PAIS = {"ES": "España", "EE": "Estonia", "SG": "Singapur", "DE": "Alemania", "PT": "Portugal"}
ASCII = {"ES": "Espana", "EE": "Estonia", "SG": "Singapur", "DE": "Alemania", "PT": "Portugal"}
ECON_KEY = {"ES": "España (ES)", "EE": "Estonia (EE)", "SG": "Singapur (SG)", "DE": "Alemania (DE)", "PT": "Portugal (PT)"}
EST = {"ok": "✅ Listo", "partial": "⚠️ Parcial / PRELIM", "manual": "❌ Falta (manual)", "na": "❌ Falta (manual)"}
FECHA = "2026-06-17"

# Banderas por país (fieles al HANDOFF §5 / .md por país).
HL = {
 "ES": "1ª agencia de IA de la UE (AESIA); 1er sandbox regulatorio de IA de la UE; EU AI Act + ley nacional; "
       "Perspectiva de género=1; GCI 5 pilares completos (20/20/20/19,74/20).",
 "EE": "Estrategia/White Paper 2024-2030 (Antigüedad=4); institucionalidad fuerte (MKM/MEAC, RIA, Eesti.ai); "
       "ISO SC 42=0 a verificar en iso.org.",
 "SG": "NAIS 2.0 + Update may-2026 (Antigüedad=4, decisión del operador); iniciativa legal=1 (soft law: marcos "
       "voluntarios muy maduros, no es debilidad).",
 "DE": "KI-Strategie 2018 → Aktionsplan 2023; €5.000M; EU AI Act + KI-MIG; GPAI fundador; Perspectiva de género=0.",
 "PT": "Nueva Agenda ANIA 2026-2030 (Antigüedad=4; >€400M); consulta pública con resultados; ISO=Observador; "
       "Género=0 y Sostenibilidad=0 en ANIA.",
}

def norm(s): return "manual" if s == "na" else s

def econ_falta(idt, status, pend):
    extra = " Ver patentes_consultas_lens_espacenet.md." if idt in ("82", "83", "84") else ""
    if status == "ok":
        return ("Validar fuente. " + (pend or "")).strip() + extra
    return (pend or "Extracción manual.") + extra

ISO_TXT = {"iso42": "Verificar membresía nacional en iso.org/committee/6794475 (SC 42).",
           "iso27": "Verificar membresía nacional en iso.org/committee/45306 (SC 27)."}
MAN_TXT = {"gci": "UIT GCI 2024: 5 pilares por país en el anexo PDF, o World Bank Data360 (indicadores ITU_GCI_*_SCORE).",
           "girai": "Portal global-index.ai → país → área (Data protection & privacy / Safety, accuracy & reliability).",
           "nri": "networkreadinessindex.org → página del país (indicador 'Affordable & clean energy')."}

def gov_falta(srckind, status, nota):
    nt = f" · Nota: {nota}" if nota else ""
    if status == "manual":
        return MAN_TXT.get(srckind, ISO_TXT.get(srckind, "Extracción manual desde la fuente del indicador.")) + nt
    if status == "partial":
        if srckind in ISO_TXT:
            return ISO_TXT[srckind] + nt
        return "Confirmar verbatim con el texto primario de la estrategia (PDF) o verificar la fuente." + nt
    return "Validar el puntaje propuesto por la rúbrica." + nt

def gid(ind):
    m = re.search(r"\((\d+)\)", ind); return m.group(1) if m else ""
def gname(ind):
    return re.sub(r"\s*\(\d+\)\s*$", "", ind)

def country_rows(cc):
    """Lista mixta de ('section'|'sub'|'row', ...) para un país."""
    rows = [("section", "A. Investigación, adopción y desarrollo (11 indicadores)")]
    cells = ECON_CELLS[ECON_KEY[cc]]
    for idt, name, sugg, desc in ECON:
        dato, status, fuentes, pend = cells[idt]
        rows.append(("row", "A", idt, name, dato, norm(status), econ_falta(idt, status, pend)))
    rows.append(("section", "B. Gobernanza (38 subindicadores)"))
    gidx = COUNTRY_GIDX[cc]
    for row in G:
        sub, ind, nota, val = row[0], row[1], row[8], row[gidx]
        if sub:
            rows.append(("sub", sub))
        idt = gid(ind); _, srckind = GOB_META.get(idt, ("", "estrategia")); s = st(val)
        rows.append(("row", "B", idt, gname(ind), val, s, gov_falta(srckind, s, nota)))
    return rows

# ---------------- 1) PLANILLA .xlsx POR PAÍS ----------------
OK = PatternFill("solid", fgColor="C6EFCE"); PART = PatternFill("solid", fgColor="FFEB9C")
MAN = PatternFill("solid", fgColor="FFC7CE"); HDR = PatternFill("solid", fgColor="1F4E78")
SEC = PatternFill("solid", fgColor="2E6FA7"); SUBF = PatternFill("solid", fgColor="DDEBF7")
WHITEB = Font(color="FFFFFF", bold=True, size=10); BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top"); THIN = Border(*[Side(style="thin", color="D9D9D9")] * 4)
FILLS = {"ok": OK, "partial": PART, "manual": MAN}

def write_xlsx(cc):
    rows = country_rows(cc)
    wb = Workbook(); ws = wb.active; ws.title = f"{PAIS[cc]} ({cc})"[:31]
    widths = [7, 40, 30, 18, 62]
    ws.cell(1, 1, f"ILIA 2026 — {PAIS[cc]} ({cc}) — estado por indicador").font = Font(bold=True, size=14, color="1F4E78")
    ws.cell(2, 1, f"Lo que hay y lo que falta por indicador. Verde=listo · Ámbar=parcial/PRELIM · Rojo=falta (manual). "
                  f"Generado {FECHA}. Detalle metodológico en data/paises_extra/.").font = Font(size=9)
    hr = 4
    for c, (h, w) in enumerate(zip(["ID", "Indicador", "Lo que hay", "Estado", "Qué falta / cómo obtenerlo"], widths), 1):
        cell = ws.cell(hr, c, h); cell.fill = HDR; cell.font = WHITEB
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[hr].height = 24
    r = hr + 1
    for row in rows:
        if row[0] == "section":
            ws.cell(r, 1, row[1]); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            ws.cell(r, 1).font = Font(bold=True, color="FFFFFF", size=11); ws.cell(r, 1).fill = SEC; r += 1; continue
        if row[0] == "sub":
            ws.cell(r, 1, row[1]); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            ws.cell(r, 1).font = BOLD; ws.cell(r, 1).fill = SUBF; r += 1; continue
        _, _blk, idt, name, dato, s, falta = row
        ws.cell(r, 1, idt); ws.cell(r, 2, name).font = BOLD; ws.cell(r, 3, dato)
        ws.cell(r, 4, EST[s]).fill = FILLS[s]; ws.cell(r, 5, falta)
        for c in range(1, 6):
            ws.cell(r, c).border = THIN; ws.cell(r, c).alignment = WRAP
        ws.row_dimensions[r].height = 30; r += 1
    ws.freeze_panes = "A5"
    out = os.path.join(HERE, "por_pais", f"Planilla_{cc}_{ASCII[cc]}_ILIA2026.xlsx")
    wb.save(out); return out

# ---------------- 2) PDF RESUMEN POR PAÍS ----------------
DJ = "/usr/share/fonts/truetype/dejavu"
pdfmetrics.registerFont(TTFont("DJ", f"{DJ}/DejaVuSans.ttf"))
pdfmetrics.registerFont(TTFont("DJ-B", f"{DJ}/DejaVuSans-Bold.ttf"))
AZUL = colors.HexColor("#1F4E78"); GRIS = colors.HexColor("#555555")
H1 = ParagraphStyle("H1", fontName="DJ-B", fontSize=15, leading=18, textColor=AZUL, spaceBefore=2, spaceAfter=2)
H2 = ParagraphStyle("H2", fontName="DJ-B", fontSize=11, leading=13, textColor=AZUL, spaceBefore=8, spaceAfter=3)
BODY = ParagraphStyle("BODY", fontName="DJ", fontSize=9, leading=11.5, spaceAfter=2)
BL = ParagraphStyle("BL", fontName="DJ", fontSize=8.7, leading=11, leftIndent=10, spaceAfter=1)
SUB = ParagraphStyle("SUB", fontName="DJ", fontSize=8.6, leading=11, textColor=GRIS, spaceAfter=3)
NOTE = ParagraphStyle("NOTE", fontName="DJ", fontSize=8.4, leading=10.5, textColor=GRIS, spaceBefore=2)

def esc(t): return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
def bullet(txt, style=BL): return Paragraph("•&nbsp;" + esc(txt), style)

def counts(rows, blk):
    c = {"ok": 0, "partial": 0, "manual": 0}
    for r in rows:
        if r[0] == "row" and r[1] == blk:
            c[r[5]] += 1
    return c

def build_pdf():
    story = [Paragraph("ILIA 2026 — Resumen por país: qué se hizo y qué queda", H1),
             Paragraph(f"5 países extra (España, Estonia, Singapur, Alemania, Portugal). Generado {FECHA}. "
                       "Detalle indicador por indicador en por_pais/Planilla_&lt;PAÍS&gt;_ILIA2026.xlsx.", SUB),
             HRFlowable(width="100%", color=AZUL, thickness=1.1, spaceAfter=4)]
    for ci, cc in enumerate(ORDER):
        if ci:
            story.append(PageBreak())
        rows = country_rows(cc)
        ec, gc = counts(rows, "A"), counts(rows, "B")
        story.append(Paragraph(f"{PAIS[cc]} ({cc})", H1))
        story.append(Paragraph(
            f"<b>Progreso —</b> Económicos: <b>{ec['ok']}/11</b> listos · {ec['partial']} parcial · {ec['manual']} por extraer.   "
            f"Gobernanza: <b>{gc['ok']}/38</b> listos · {gc['partial']} por confirmar · {gc['manual']} por extraer.", BODY))
        story.append(Paragraph(f"<b>Banderas:</b> {esc(HL[cc])}", NOTE))

        # Qué se hizo
        story.append(Paragraph("Qué se hizo", H2))
        hechos = [r for r in rows if r[0] == "row" and r[5] == "ok"]
        econ_ok = [r for r in hechos if r[1] == "A"]
        if econ_ok:
            story.append(Paragraph("<b>Económicos con dato:</b>", BODY))
            for _, _b, idt, name, dato, s, falta in econ_ok:
                story.append(bullet(f"[{idt}] {name}: {dato}"))
        story.append(Paragraph(f"<b>Gobernanza:</b> {gc['ok']}/38 subindicadores con puntaje propuesto por la rúbrica "
                               f"(verde), listos para validar.", BODY))

        # Qué queda
        story.append(Paragraph("Qué queda (acción)", H2))
        econ_pend = [r for r in rows if r[0] == "row" and r[1] == "A" and r[5] != "ok"]
        if econ_pend:
            story.append(Paragraph("<b>Económicos por obtener/validar:</b>", BODY))
            for _, _b, idt, name, dato, s, falta in econ_pend:
                story.append(bullet(f"[{idt}] {name} — {EST[s]} · {falta}"))
        gov_man = [r for r in rows if r[0] == "row" and r[1] == "B" and r[5] == "manual"]
        if gov_man:
            story.append(Paragraph("<b>Gobernanza por extraer (manual):</b>", BODY))
            for _, _b, idt, name, dato, s, falta in gov_man:
                story.append(bullet(f"[{idt}] {name} — {falta}"))
        gov_par = [r for r in rows if r[0] == "row" and r[1] == "B" and r[5] == "partial"]
        if gov_par:
            nombres = ", ".join(f"[{idt}] {name} ({val})" for _, _b, idt, name, val, s, falta in gov_par)
            story.append(Paragraph("<b>Gobernanza por confirmar (PRELIM/duda):</b> " + esc(nombres) +
                                   ". Confirmar con el texto primario de la estrategia (PDF) o la fuente.", BODY))
        # insumos/decisiones pendientes
        ins = ["Patentes (82-84): definición CPC + ventana = replicar ILIA 2025 (falta su especificación)."]
        if cc == "EE":
            ins.append("ISO SC 42 (125): verificación manual en iso.org (hoy 0, duda).")
        story.append(Paragraph("<b>Insumos/decisiones pendientes:</b> " + esc(" ".join(ins)), NOTE))
        story.append(Paragraph(f"Detalle indicador por indicador: <b>por_pais/Planilla_{cc}_{ASCII[cc]}_ILIA2026.xlsx</b>.", NOTE))

    out = os.path.join(HERE, "Resumen_por_pais_ILIA2026.pdf")
    SimpleDocTemplate(out, pagesize=A4, topMargin=1.1 * cm, bottomMargin=1.1 * cm,
                      leftMargin=1.4 * cm, rightMargin=1.4 * cm,
                      title="ILIA 2026 — Resumen por país").build(story)
    return out

if __name__ == "__main__":
    xs = [write_xlsx(cc) for cc in ORDER]
    for x in xs:
        print("OK xlsx ->", x)
    print("OK pdf  ->", build_pdf())
