#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Planilla orientada a PROGRAMA (puede haber >1 por universidad) para revision de Nicole.
Fuente: clasificacion.csv (1 programa base/univ) + EXTRAS documentados + URLs scrapeadas
de paises/*.md (nivel universidad, best-effort).
Salida: entregables/Programas_por_Universidad.xlsx  (+ .csv)
"""
import csv, os, re, datetime
from collections import OrderedDict, defaultdict

BASE = os.path.dirname(os.path.abspath(__file__))
OUT  = os.path.join(BASE, "entregables"); os.makedirs(OUT, exist_ok=True)
PAIS_ORDER = ["Alemania","España","Portugal","Singapur","Estonia"]
PAIS_MD = {"Alemania":"alemania.md","España":"espana.md","Portugal":"portugal.md",
           "Singapur":"singapur.md","Estonia":"estonia.md"}

def read_csv(p):
    with open(p, encoding="utf-8") as f: return list(csv.DictReader(f))
clas = read_csv(os.path.join(BASE,"clasificacion.csv"))

# ---- EXTRAS: programas adicionales documentados (univ debe coincidir con clasificacion.csv) ----
EXTRAS = [
 # pais, universidad, programa, facultad, clasificacion, url
 ("Portugal","Universidade de Lisboa (Inst. Superior Tecnico)","PhD in Electrical and Computer Engineering","IST (INESC-ID/INESC)","CUMPLE","https://tecnico.ulisboa.pt/en/education/courses/phd-programmes/"),
 ("Portugal","Universidade de Lisboa (Inst. Superior Tecnico)","Doutoramento em Informatica (Fac. Ciencias - FCUL/LASIGE)","FCUL (LASIGE)","CUMPLE","https://ciencias.ulisboa.pt/"),
 ("Portugal","Universidade NOVA de Lisboa (NOVA FCT)","PhD in Information Management (Data Science) - NOVA IMS","NOVA IMS","CUMPLE","https://www.novaims.unl.pt/en/education/programs/doctoral-program-in-information-management/"),
 ("España","Universidad de Salamanca (USAL)","Doctorado en Ingenieria Informatica (ademas del de IA Aplicada)","Esc. Doctorado (GRIAL/BISITE)","CUMPLE","https://doctorado.usal.es/"),
 ("Alemania","University of Tubingen (Eberhard Karls)","IMPRS for Intelligent Systems (escuela doctoral estructurada de IA)","Tubingen AI Center / MPI-IS","CUMPLE","https://imprs.is.mpg.de/"),
 ("Alemania","TU Dresden","ScaDS.AI Graduate School (ademas del Promotion CS)","ScaDS.AI Dresden-Leipzig","CUMPLE","https://scads.ai/education/graduate-school/"),
]

# ---- Scrape URLs por universidad desde los .md (best-effort, nivel universidad) ----
URL_RE = re.compile(r'https?://[^\s\)\]\`<>"]+')
BT_RE  = re.compile(r'`([^`]+)`')
DOM_RE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9.\-]+\.[A-Za-z]{2,}(/\S*)?$')
STOP = set(("university universidad universidade universitat universitaet of de la the and y e technical "
            "institute institut national autonoma politecnica catolica madrid berlin barcelona munich "
            "munchen muenchen valencia catalunya lisboa lisbon london paris science technology computer "
            "informatica information rey juan carlos santiago compostela gran canaria").lower().split())

HEADER_RE = re.compile(r'^\s{0,3}(#{1,6}\s|\||\*\*|\d+\s*[\./])')
def is_header(line):
    return bool(HEADER_RE.match(line)) or (' — ' in line)

def urls_in(line):
    out=set()
    for u in URL_RE.findall(line): out.add(u.rstrip('.,);*`'))
    for seg in BT_RE.findall(line):
        s=seg.strip()
        if s.lower().startswith("http"): out.add(s.rstrip('.,);*`'))
        elif DOM_RE.match(s): out.add("https://"+s.rstrip('.,);*`'))
    return out

def aliases_for(name):
    al=[]
    for p in re.findall(r'\(([^)]+)\)', name):
        p=p.strip()
        if 2<=len(p)<=24: al.append(p)
    base=re.sub(r'\([^)]*\)','',name)
    words=[w for w in re.split(r'[\s/,-]+', base) if len(w)>=4 and w.lower() not in STOP]
    words.sort(key=len, reverse=True)
    al += words[:2]
    # normaliza, mas largos primero, sin duplicados
    seen=set(); res=[]
    for a in sorted(al, key=len, reverse=True):
        k=a.lower()
        if k not in seen: seen.add(k); res.append(a)
    return res

url_by_univ = defaultdict(set)
for pais in PAIS_ORDER:
    path=os.path.join(BASE,"paises",PAIS_MD[pais])
    univs=[r["universidad"] for r in clas if r["pais"]==pais]
    al_map=[(u, aliases_for(u)) for u in univs]
    cur=None
    for line in open(path, encoding="utf-8"):
        low=line.lower()
        # cambia universidad actual SOLO en lineas de cabecera/tabla (evita saltos por palabras sueltas)
        if is_header(line):
            best=None; best_len=0
            for u,als in al_map:
                for a in als:
                    if a.lower() in low and len(a)>best_len:
                        best=u; best_len=len(a)
            if best is not None:
                cur=best
        if cur is not None:
            for u in urls_in(line): url_by_univ[cur].add(u)

# ---- Construye filas-programa ----
COLS=["pais","universidad","qs_rank_2026","programa_doctorado","facultad_escuela",
      "clasificacion","url_programa","fuente_origen","para_revision"]
prog_rows=[]
for r in clas:
    prog_rows.append({"pais":r["pais"],"universidad":r["universidad"],"qs_rank_2026":r["qs_rank_2026"],
        "programa_doctorado":r["programa_doctorado"],"facultad_escuela":r["facultad_escuela"],
        "clasificacion":r["clasificacion"],"url_programa":r["fuente_principal"],
        "fuente_origen":"clasificacion (programa base)","para_revision":""})
for (pais,uni,prog,fac,cl,url) in EXTRAS:
    prog_rows.append({"pais":pais,"universidad":uni,"qs_rank_2026":next((r["qs_rank_2026"] for r in clas if r["universidad"]==uni),""),
        "programa_doctorado":prog,"facultad_escuela":fac,"clasificacion":cl,"url_programa":url,
        "fuente_origen":"EXTRA documentado","para_revision":"programa adicional"})

# orden por pais y universidad
prog_rows.sort(key=lambda x:(PAIS_ORDER.index(x["pais"]), x["universidad"]))

# nº programas por universidad
nprog=defaultdict(int)
for r in prog_rows: nprog[(r["pais"],r["universidad"])]+=1

# CSV
csv_path=os.path.join(OUT,"Programas_por_Universidad.csv")
with open(csv_path,"w",encoding="utf-8",newline="") as f:
    w=csv.DictWriter(f, fieldnames=COLS); w.writeheader()
    for r in prog_rows: w.writerow(r)

# ---- XLSX ----
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
GREEN=PatternFill("solid",fgColor="C6EFCE"); YELLOW=PatternFill("solid",fgColor="FFEB9C"); RED=PatternFill("solid",fgColor="FFC7CE")
HEAD=PatternFill("solid",fgColor="305496"); HF=Font(bold=True,color="FFFFFF"); THIN=Border(*[Side(style="thin",color="D9D9D9")]*4)
def fill(c): return {"CUMPLE":GREEN,"PARCIAL":YELLOW,"NO_CUMPLE":RED}.get(c)
wb=Workbook()

# Hoja Programas
ws=wb.active; ws.title="Programas"
hdr=["Pais","Universidad","QS 2026","N programas (univ)","Programa de doctorado","Facultad/Escuela","Clasificacion","URL del programa","Origen","Para revision"]
for j,h in enumerate(hdr,1):
    c=ws.cell(1,j,h); c.fill=HEAD; c.font=HF; c.alignment=Alignment(horizontal="center",wrap_text=True)
for i,r in enumerate(prog_rows,2):
    vals=[r["pais"],r["universidad"],r["qs_rank_2026"],nprog[(r["pais"],r["universidad"])],r["programa_doctorado"],
          r["facultad_escuela"],r["clasificacion"],r["url_programa"],r["fuente_origen"],r["para_revision"]]
    for j,v in enumerate(vals,1):
        c=ws.cell(i,j,v); c.border=THIN; c.alignment=Alignment(vertical="top",wrap_text=True)
    f=fill(r["clasificacion"]);
    if f: ws.cell(i,7).fill=f
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(hdr))}{len(prog_rows)+1}"
W=[10,34,9,11,42,30,12,46,22,28]
for j,w_ in enumerate(W,1): ws.column_dimensions[get_column_letter(j)].width=w_

# Hoja Resumen_universidades
ws2=wb.create_sheet("Resumen_universidades")
h2=["Pais","Universidad","QS 2026","N programas","Programas (lista)","Clasificacion(es)","Todas las URLs recopiladas (best-effort)"]
for j,h in enumerate(h2,1):
    c=ws2.cell(1,j,h); c.fill=HEAD; c.font=HF; c.alignment=Alignment(horizontal="center",wrap_text=True)
univ_order=[]
for r in prog_rows:
    k=(r["pais"],r["universidad"])
    if k not in univ_order: univ_order.append(k)
for i,(pais,uni) in enumerate(univ_order,2):
    progs=[r for r in prog_rows if r["pais"]==pais and r["universidad"]==uni]
    qs=progs[0]["qs_rank_2026"]
    plist=" | ".join(p["programa_doctorado"] for p in progs)
    cls=" / ".join(sorted(set(p["clasificacion"] for p in progs)))
    allurls=set(p["url_programa"] for p in progs) | url_by_univ.get(uni,set())
    urls="  ".join(sorted(u for u in allurls if u))
    vals=[pais,uni,qs,len(progs),plist,cls,urls]
    for j,v in enumerate(vals,1):
        c=ws2.cell(i,j,v); c.border=THIN; c.alignment=Alignment(vertical="top",wrap_text=True)
    if len(progs)>1: ws2.cell(i,4).fill=PatternFill("solid",fgColor="FFF2CC")
ws2.freeze_panes="A2"; ws2.auto_filter.ref=f"A1:G{len(univ_order)+1}"
for j,w_ in enumerate([10,34,9,11,52,16,70],1): ws2.column_dimensions[get_column_letter(j)].width=w_

# Hoja Leeme
ws3=wb.create_sheet("Leeme"); ws3.column_dimensions["A"].width=110
txt=[
 ("Planilla de PROGRAMAS de doctorado de IA por universidad (QS 2026) - para revision",True),
 (f"Generado {datetime.date.today().isoformat()}. Universidades: {len(univ_order)}. Programas listados: {len(prog_rows)}.",False),
 ("",False),
 ("COMO LEERLA:",True),
 ("- Hoja 'Programas': UNA FILA POR PROGRAMA. Una universidad puede aparecer en varias filas.",False),
 ("- Hoja 'Resumen_universidades': una fila por universidad, con N de programas y todas las URLs recopiladas.",False),
 ("- Color en 'Clasificacion': verde=CUMPLE, amarillo=PARCIAL, rojo=NO_CUMPLE.",False),
 ("",False),
 ("IMPORTANTE - el N de programas es un MINIMO IDENTIFICADO, no exhaustivo:",True),
 ("- La investigacion se hizo SOLO con busqueda web (sin poder abrir paginas), por lo que para la mayoria de",False),
 ("  universidades se identifico 1 programa representativo. Se anadieron como 'EXTRA documentado' los programas",False),
 ("  adicionales con evidencia clara. NICOLE: completar con los demas programas Tec/Ing con IA de cada universidad",False),
 ("  (anadir filas en la hoja 'Programas').",False),
 ("- 'URL del programa' = fuente principal de ese programa. 'Todas las URLs recopiladas' (Resumen) incluye el resto",False),
 ("  de enlaces que aparecen en paises/*.md (atribuidos a la universidad de forma best-effort: verificar).",False),
 ("- Para la evidencia y todas las fuentes por universidad, ver los archivos paises/*.md del repositorio.",False),
 ("",False),
 ("Las cifras del indicador (CUMPLE/PARCIAL/NO_CUMPLE) estan en RESUMEN.md y en Planilla_Doctorados_IA_QS2026.xlsx.",False),
]
for i,(t,b) in enumerate(txt,1):
    c=ws3.cell(i,1,t); c.font=Font(bold=b); c.alignment=Alignment(wrap_text=True,vertical="top")

xlsx=os.path.join(OUT,"Programas_por_Universidad.xlsx"); wb.save(xlsx)

# ---- reporte ----
multi=[(p,u,nprog[(p,u)]) for (p,u) in univ_order if nprog[(p,u)]>1]
print(f"Universidades: {len(univ_order)} | Programas listados: {len(prog_rows)} | con >1 programa: {len(multi)}")
for p,u,n in multi: print(f"  {p:9s} {u}  -> {n}")
print("Spot-check URLs recopiladas:")
for u in ["University of Tartu","Universitat Politecnica de Catalunya (UPC)","Technical University of Munich (TUM)"]:
    print(f"  {u}: {len(url_by_univ.get(u,set()))} urls")
print("XLSX:",xlsx); print("CSV:",csv_path)
